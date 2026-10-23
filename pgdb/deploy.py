"""Provision the DST & TT database from nothing, and prove it works.

One command takes an empty PostgreSQL database to a loaded, tested one:
schema, reference data, views, roles, the courier and ROTN migrations, the
sales history, the targets, then the self-test.

    set DATABASE_URL=postgresql://user:pw@host:5432/dstt
    python deploy.py --excel "C:\\path\\DST & TT DATABASE.xlsx"

    python deploy.py --dry-run          # say what it would do, touch nothing
    python deploy.py --schema-only      # structure and reference data, no rows
    python deploy.py --upgrade          # migrations only, preserves existing data

Fresh provisioning refuses an existing schema. --upgrade applies only
repeatable migrations, never the initial CREATE TABLE statements or loaders.

Needs psql for nothing -- everything goes through psycopg2.
"""

import argparse
import csv
import os
import subprocess
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
SQL = os.path.join(HERE, "sql")
ETL = os.path.join(HERE, "etl")

try:
    import psycopg2
except ImportError:
    sys.exit("psycopg2 is not installed.  pip install -r etl/requirements.txt")


# Order matters. 05 and 06 are migrations for databases built before those
# features existed; on a fresh install 01 already carries their tables, and
# both are written to be harmless when re-run.
SCHEMA_FILES = ["01_schema.sql", "02_seed.sql", "03_views.sql", "04_security.sql"]
MIGRATIONS   = ["05_add_courier.sql", "07_web_entry.sql", "08_auto_entry_numbers.sql",
                "09_supplementary_data.sql", "10_sheet_sales_sync.sql", "11_lfo_secondary.sql", "12_lfo_secondary_targets.sql", "13_lfo_branch_targets.sql",
                "14_llfo_secondary.sql", "15_target_planner.sql", "16_shipping_details.sql", "17_product_invoice.sql", "18_invoice_colour.sql", "19_online_sales.sql"]
POST_LOAD    = ["06_add_rotn.sql"]      # needs the people the loader creates
UPGRADE_FILES = ["04_security.sql", "05_add_courier.sql", "06_add_rotn.sql",
                 "07_web_entry.sql", "08_auto_entry_numbers.sql", "09_supplementary_data.sql", "10_sheet_sales_sync.sql", "11_lfo_secondary.sql", "12_lfo_secondary_targets.sql", "13_lfo_branch_targets.sql",
                 "14_llfo_secondary.sql", "15_target_planner.sql", "16_shipping_details.sql", "17_product_invoice.sql", "18_invoice_colour.sql", "19_online_sales.sql", "03_views.sql"]


def say(msg=""):
    print(msg, flush=True)


def step(n, total, title):
    say(f"\n[{n}/{total}] {title}")
    say("-" * 66)


def run_sql(dsn, path, dry):
    name = os.path.basename(path)
    if dry:
        say(f"    would run {name}")
        return
    with open(path, encoding="utf-8") as f:
        body = f.read()
    conn = psycopg2.connect(dsn)
    conn.autocommit = True          # the files carry their own begin/commit
    try:
        with conn.cursor() as cur:
            cur.execute(body)       # no args, so % inside the SQL is left alone
        say(f"    ok  {name}")
    except psycopg2.Error as e:
        say(f"    FAILED  {name}")
        say(f"            {str(e).strip()}")
        raise SystemExit(1)
    finally:
        conn.close()


def run_py(script, args, dry, env):
    if dry:
        say(f"    would run {script} {' '.join(args)}")
        return ""
    cmd = [sys.executable, os.path.join(ETL, script)] + args
    p = subprocess.run(cmd, cwd=ETL, env=env, capture_output=True, text=True)
    out = (p.stdout or "") + (p.stderr or "")
    for line in out.splitlines():
        if line.strip():
            say("    " + line.rstrip())
    if p.returncode != 0:
        raise SystemExit(f"{script} failed with exit code {p.returncode}")
    return out


def schema_present(dsn):
    conn = psycopg2.connect(dsn)
    try:
        with conn.cursor() as cur:
            cur.execute("select count(*) from information_schema.tables"
                        " where table_schema in ('master','txn','tgt','sec')")
            return cur.fetchone()[0] > 0
    finally:
        conn.close()


def preflight(args):
    """Validate local inputs before opening a write connection."""
    if args.force:
        sys.exit("--force is no longer supported. Use --upgrade for an existing "
                 "database; use a new empty database for a fresh import.")
    if args.upgrade and (args.schema_only or args.excel):
        sys.exit("--upgrade cannot be combined with --schema-only or --excel. "
                 "Run the loaders separately after checking their dry-run reports.")
    files = UPGRADE_FILES if args.upgrade else SCHEMA_FILES + MIGRATIONS + POST_LOAD
    for name in files:
        if not os.path.isfile(os.path.join(SQL, name)):
            sys.exit(f"missing SQL migration: {name}")
    if args.schema_only or args.upgrade:
        return
    if not args.excel:
        sys.exit("--excel is required unless you pass --schema-only or --upgrade")
    for label, path in (("workbook", args.excel), ("target CSV", args.targets)):
        if not os.path.isfile(path):
            sys.exit(f"{label} not found: {path}. No schema changes were made.")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(args.excel, read_only=True, data_only=True)
        try:
            if "Direct Sales Report" not in wb.sheetnames:
                sys.exit("workbook must contain a 'Direct Sales Report' sheet")
            header = {str(v).strip() for v in next(
                wb["Direct Sales Report"].iter_rows(values_only=True), ()) if v is not None}
            for choices in (("Tax Invoice Date", "Invoice Date"),
                            ("Supplier Name", "Customer Name"), ("Channels", "Region"),
                            ("DT / TT",), ("DT/TT Primary Taxable Value ( Without Tax)",)):
                if not any(name in header for name in choices):
                    sys.exit(f"workbook is missing a required column: {choices[0]}")
        finally:
            wb.close()
        with open(args.targets, newline="", encoding="utf-8-sig") as f:
            header = {v.strip().lower() for v in (csv.DictReader(f).fieldnames or [])}
            if not {"period_month", "target_value"} <= header:
                sys.exit("target CSV needs period_month and target_value columns")
    except (OSError, ValueError, ImportError) as exc:
        sys.exit(f"input preflight failed: {exc}. No schema changes were made.")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dsn", default=os.environ.get("DATABASE_URL"),
                    help="postgresql://... (default: $DATABASE_URL)")
    ap.add_argument("--excel", help="the sales workbook to load")
    ap.add_argument("--targets", default=os.path.join(ETL, "targets.example.csv"),
                    help="target CSV (default: fictional example data; supply your own file for real use)")
    ap.add_argument("--fy", type=int, default=2026)
    ap.add_argument("--plan", default="FY26-27 Base")
    ap.add_argument("--schema-only", action="store_true", help="no data")
    ap.add_argument("--force", action="store_true",
                    help=argparse.SUPPRESS)
    ap.add_argument("--upgrade", action="store_true",
                    help="apply repeatable migrations to an existing database; no import")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    preflight(args)

    if not args.dsn:
        sys.exit("no DATABASE_URL set and no --dsn given")

    # Never print the password back at the operator.
    safe = args.dsn
    if "@" in safe:
        safe = safe.split("://")[0] + "://" + safe.split("@", 1)[1]

    say("\n" + "=" * 66)
    say("  DST & TT  --  database deployment")
    say("=" * 66)
    say(f"  target        {safe}")
    say(f"  mode          {'DRY RUN (nothing is written)' if args.dry_run else 'live'}")

    t0 = time.time()

    # ---------------------------------------------------------- 1. connect
    total = 5 if args.schema_only else 6
    step(1, total, "Connect and check what is already there")
    try:
        conn = psycopg2.connect(args.dsn)
        with conn.cursor() as cur:
            cur.execute("select version(), current_database(), current_user")
            ver, db, who = cur.fetchone()
        conn.close()
    except psycopg2.Error as e:
        sys.exit(f"    cannot connect: {str(e).strip()}")
    say(f"    {ver.split(',')[0]}")
    say(f"    database '{db}' as '{who}'")

    present = schema_present(args.dsn)
    if args.upgrade:
        if not present:
            sys.exit("--upgrade requires an existing schema. Use --schema-only or --excel first.")
        say("    upgrading existing database; no sales or targets will be imported")
        for f in UPGRADE_FILES:
            run_sql(args.dsn, os.path.join(SQL, f), args.dry_run)
        say("\n    Upgrade complete." if not args.dry_run else "\n    Upgrade dry run complete.")
        say("    Verify application-role access and run pgdb/selftest.py on a test copy.")
        return
    if present:
        sys.exit("\n    This database already has the schema. Use --upgrade for migrations,\n"
                 "    or use an empty database for a fresh import.")
    else:
        say("    empty database")

    # ----------------------------------------------------------- 2. schema
    step(2, total, "Schema, reference data, views and roles")
    for f in SCHEMA_FILES:
        run_sql(args.dsn, os.path.join(SQL, f), args.dry_run)

    step(3, total, "Migrations")
    for f in MIGRATIONS:
        run_sql(args.dsn, os.path.join(SQL, f), args.dry_run)

    env = dict(os.environ, DATABASE_URL=args.dsn)

    if args.schema_only:
        step(4, total, "Data load -- skipped (--schema-only)")
    else:
        # ------------------------------------------------------- 4. the data
        step(4, total, "Load the sales history and the targets")
        run_py("load_sales.py", [args.excel], args.dry_run, env)
        run_py("load_targets.py", [args.targets, "--fy", str(args.fy),
                                   "--plan", args.plan, "--activate"],
               args.dry_run, env)

    # ------------------------------------------- 5. the post-load migration
    step(5, total, "Territory map (needs the people the loader created)")
    for f in POST_LOAD:
        run_sql(args.dsn, os.path.join(SQL, f), args.dry_run)
    # 06 adds a column the views select, so they have to be rebuilt after it.
    run_sql(args.dsn, os.path.join(SQL, "03_views.sql"), args.dry_run)

    # --------------------------------------------------------- 6. self-test
    if not args.schema_only:
        step(6, total, "Self-test")
        if args.dry_run:
            say("    would run selftest.py")
        else:
            p = subprocess.run([sys.executable, os.path.join(HERE, "selftest.py")],
                               cwd=HERE, env=env, capture_output=True, text=True)
            for line in ((p.stdout or "") + (p.stderr or "")).splitlines():
                if line.strip():
                    say("    " + line.rstrip())
            if p.returncode != 0:
                sys.exit("\n    THE SELF-TEST FAILED. The database is not fit to serve.")

    say("\n" + "=" * 66)
    if args.dry_run:
        say(f"  Dry run complete in {time.time() - t0:.1f}s. Nothing was written.")
    else:
        say(f"  Deployed in {time.time() - t0:.1f}s.")
        say("")
        say("  Next:")
        say("    1. New database roles are NOLOGIN with no password. Enable LOGIN")
        say("       and set separate strong credentials with your provider or psql")
        say("       \\password. Existing role credentials have not been changed.")
        say("    2. Put the connection string on the app server:")
        say("         powershell -File setup-server-task.ps1 -DatabaseUrl '...'")
        say("    3. Run compare.py for a month before switching the sheet off.")
    say("=" * 66 + "\n")


if __name__ == "__main__":
    main()
