"""Export the current dashboard's workbook targets for load_targets.py.

This reads a saved primary workbook only; it never fetches Sheets or connects to
PostgreSQL. The dashboard's parsers determine years, months and rounded values.
Derived totals are omitted so channel targets are counted once. Executive
targets retain the person/place scope supported by pgdata.fetch_exec_targets.

Example:
    python pgdb/etl/export_dashboard_targets.py current.xlsx --fy 2026 \
        --output current-targets.csv --report current-targets.json

An unresolved/ambiguous place rejects the entire export. Supply --place-map
with a JSON object such as {"Trichy": {"state": "Tamil Nadu", "city": "Trichy"}}
only after verifying the mapping. A state-wide target uses an empty city.
"""

import argparse
import csv
import hashlib
import json
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import refresh
from pgdb.etl.load_sales import CITY_FIX, CITY_STATE_FIX, DEFAULT_ZONE, STATE_FIX


FIELDS = ["period_month", "metric", "channel", "sub_channel", "person", "state",
          "district", "city", "brand", "target_value", "is_rollup", "notes"]
CHANNEL_CODES = {value: key for key, value in refresh.CHANNEL_TABS.items()}


class ExportProblem(ValueError):
    pass


def clean(value):
    return str(value or "").strip()


def strict_value(value, where):
    """Reject malformed cells instead of letting the dashboard coerce them to 0."""
    if value is None or value == "":
        return
    try:
        amount = Decimal(str(value).replace(",", "").replace("₹", "").strip())
    except InvalidOperation:
        raise ExportProblem(f"{where}: unreadable target amount {value!r}") from None
    if not amount.is_finite() or not 0 <= amount < Decimal("1000000000000"):
        raise ExportProblem(f"{where}: target amount must be finite, nonnegative and fit numeric(14,2)")


def validate_source(wb, fy):
    """A dashboard last-row-wins fallback must not hide a migration conflict."""
    for tab in refresh.CHANNEL_TABS:
        if tab not in wb.sheetnames:
            raise ExportProblem(f"Missing channel target tab: {tab}")
        ws = wb[tab]
        hdr = [refresh.norm(c.value) for c in ws[1]]
        if "Month" not in hdr or "Target" not in hdr:
            raise ExportProblem(f"{tab}: Month and Target headers are required")
        mi, ti, yi = hdr.index("Month"), hdr.index("Target"), refresh.year_col(hdr)
        seen = set()
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if row[mi] is None:
                continue
            month, year = refresh.month_year(row[mi], fy)
            if month is None:
                if row[ti] not in (None, ""):
                    raise ExportProblem(f"{tab}!{n}: unreadable month with a target: {row[mi]!r}")
                continue
            if yi is not None and yi < len(row) and row[yi] is not None:
                year = refresh.fy_of(f"{month}-{int(refresh.num(row[yi]))}") or year
            scope = year, month
            if scope in seen:
                raise ExportProblem(f"{tab}!{n}: repeated month {month}, FY {year}; resolve before export")
            seen.add(scope)
            strict_value(row[ti], f"{tab}!{n}")
    if refresh.EXEC_TAB not in wb.sheetnames:
        raise ExportProblem("Executive tab is missing; migration cannot silently discard cached targets")
    ws = wb[refresh.EXEC_TAB]
    month_columns = []
    seen = set()
    for j, cell in enumerate(ws[1]):
        month, year = refresh.month_year(cell.value, fy)
        if month is not None:
            if (year, month) in seen:
                raise ExportProblem(f"Executive: repeated {month}, FY {year} column")
            seen.add((year, month))
            month_columns.append(j)
    if not month_columns:
        raise ExportProblem("Executive tab has no readable month columns")
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not clean(row[0]):
            continue
        for j in month_columns:
            strict_value(row[j], f"Executive!{n}, column {j + 1}")


def geography(wb):
    """Use exactly the loader's spelling repairs, retaining all state candidates."""
    ws = wb[refresh.LEDGER_TAB]
    hdr = [refresh.norm(c.value) for c in ws[1]]
    if "State" not in hdr or "City" not in hdr:
        raise ExportProblem("Sales ledger needs State and City columns to resolve target places")
    si, ci = hdr.index("State"), hdr.index("City")
    states = {name.casefold(): name for name in DEFAULT_ZONE}
    cities = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        state, city = clean(row[si]), clean(row[ci])
        state = STATE_FIX.get(state.casefold(), state)
        city = CITY_FIX.get(city.casefold(), city)
        state = CITY_STATE_FIX.get(city.casefold(), state)
        if state and state.casefold() != "unknown":
            states.setdefault(state.casefold(), state)
            if city and city.casefold() != "unknown":
                cities[city.casefold()].add((states[state.casefold()], city))
    return states, cities


def resolve_place(place, states, cities, place_map):
    if not place:
        raise ExportProblem("Executive target has a blank place; refusing a company-wide scope")
    if place in place_map:
        resolved = place_map[place]
        if not isinstance(resolved, dict) or set(resolved) != {"state", "city"}:
            raise ExportProblem(f"Explicit map for {place!r} must contain exactly state and city")
        state, city = clean(resolved["state"]), clean(resolved["city"])
        if not state:
            raise ExportProblem(f"Explicit map for {place!r} requires a state")
        return state, city, "explicit place map"
    key = place.casefold()
    candidate_city = CITY_FIX.get(key, place)
    matches = cities.get(candidate_city.casefold(), set())
    if key in states and matches:
        raise ExportProblem(f"{place!r} names both a city and state; supply --place-map")
    if key in states:
        return states[key], "", "state name"
    if len(matches) == 1:
        state, city = next(iter(matches))
        return state, city, "unique city/state in sales ledger"
    raise ExportProblem(f"Cannot uniquely place {place!r}: {sorted(matches)}; supply --place-map")


def export_rows(wb, fy, place_map=None):
    validate_source(wb, fy)
    channel_by_fy = refresh.parse_channel_targets(wb, fy)
    executive_by_fy = refresh.parse_exec_targets(wb, fy)
    channels = channel_by_fy.get(fy, {})
    executives = executive_by_fy.get(fy, [])
    if not channels:
        raise ExportProblem(f"No channel targets for FY {fy}")
    states, cities = geography(wb)
    rows, mappings, seen = [], [], set()

    def append(month, value, **dimensions):
        if not 0 <= value < 1000000000000:
            raise ExportProblem(f"Rounded target for {month} does not fit numeric(14,2)")
        number = refresh.MN.index(month) + 1
        year = fy if number >= 4 else fy + 1
        row = dict.fromkeys(FIELDS, "")
        row.update(period_month=f"{year:04d}-{number:02d}-01", metric="primary_taxable",
                   target_value=value, **dimensions)
        scope = tuple(str(row[k]).casefold() for k in FIELDS[:9])
        if scope in seen:
            raise ExportProblem(f"Duplicate target scope: {row['person']} / {row['city'] or row['state']} / {month}")
        seen.add(scope)
        rows.append(row)

    for key, code in CHANNEL_CODES.items():
        for month, value in channels.get(key, {}).items():
            append(month, value, channel=code, is_rollup="Y", notes=f"Current workbook: {code} / {month}")
    for executive in executives:
        state, city, evidence = resolve_place(executive["city"], states, cities, place_map or {})
        mappings.append(dict(person=executive["person"], source_place=executive["city"],
                             state=state, city=city, evidence=evidence))
        for month, value in executive["full"].items():
            append(month, value, person=executive["person"], state=state, city=city,
                   is_rollup="N", notes=f"Current workbook: Executive / {executive['city']} / {month}")
    report = {
        "financial_year": fy,
        "row_count": len(rows),
        "channel_rows": sum(row["is_rollup"] == "Y" for row in rows),
        "executive_rows": sum(row["is_rollup"] == "N" for row in rows),
        "company_monthly_target": channels["ALL TOTAL"],
        "company_annual_target": sum(channels["ALL TOTAL"].values()),
        "source_dashboard_channel_targets": channels,
        "source_dashboard_executive_targets": executives,
        "place_mappings": mappings,
        "expected_import_rows": rows,
        "source_target_financial_years": sorted(set(channel_by_fy) | set(executive_by_fy)),
        "notes": ["Amounts follow the current dashboard's whole-rupee rounding.",
                  "Only the five channel scopes are company rollups; executive targets overlap.",
                  "Executive targets use person and city/state with no channel, as required by the existing dashboard database reader.",
                  "The source Executive parser fills missing month amounts with zero; those explicit zeros are retained."]
    }
    try:
        report["source_dashboard_rotn"] = refresh.parse_rotn(wb, fy)
    except refresh.TabProblem as exc:
        report["source_dashboard_rotn"] = None
        report["notes"].append(f"ROTN mapping could not be read: {exc}")
    used_tabs = {*refresh.CHANNEL_TABS, refresh.EXEC_TAB, refresh.LEDGER_TAB, refresh.ROTN_TAB}
    report["other_tabs_not_used_by_current_primary_dashboard_parsers"] = [
        tab for tab in wb.sheetnames if tab not in used_tabs]
    return rows, report


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--fy", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--place-map", type=Path)
    args = parser.parse_args()
    paths = [args.workbook.resolve(), args.output.resolve(), args.report.resolve()]
    if len(set(paths)) != len(paths):
        parser.error("workbook, output CSV and report must be different paths")
    try:
        place_map = json.loads(args.place_map.read_text(encoding="utf-8-sig")) if args.place_map else {}
        # Match refresh.fetch's normal worksheet mode: Sheets exports may have
        # incomplete dimension metadata, which makes read_only yield empty tuples.
        wb = openpyxl.load_workbook(args.workbook, data_only=True)
        try:
            rows, report = export_rows(wb, args.fy, place_map)
        finally:
            wb.close()
        report["source_workbook"] = str(args.workbook.resolve())
        report["source_sha256"] = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
        with args.output.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(rows)
        args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    except (ExportProblem, refresh.TabProblem, OSError, ValueError) as exc:
        parser.exit(1, f"Target export rejected: {exc}\n")
    print(f"Exported {len(rows)} targets for FY {args.fy}; company annual target {report['company_annual_target']:,}.")
    print(f"CSV: {args.output}\nReconciliation: {args.report}")


if __name__ == "__main__":
    main()
