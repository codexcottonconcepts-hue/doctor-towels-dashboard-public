"""Import explicitly mapped Executive worksheet targets into an existing plan.

Defaults to a read-only preview. --apply inserts only absent scopes; different
existing values or overlapping person scopes reject the whole operation. This
never activates a plan, edits a channel total, deletes a target, or creates a
master record. Blank source cells remain unset; entered zero is retained.

Example mapping file (each assignment must explicitly name its channel):
  [{"person":"Demo Executive","place":"Andhra Pradesh","channel":"D&R-S",
    "state":"Andhra Pradesh","city":""}]

Use --expect-manifest with --apply to bind the write to the reviewed preview.
"""
import argparse
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import tempfile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
import pgdata


class ImportProblem(ValueError):
    pass


MONTHS = ('Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec', 'Jan', 'Feb', 'Mar')
ASSIGNMENT_KEYS = {'person', 'place', 'channel', 'state', 'city'}


def clean(value):
    return '' if value is None else str(value).strip()


def amount(value, where):
    if value is None or clean(value) == '':
        return None
    if isinstance(value, bool):
        raise ImportProblem(f'{where}: boolean is not a target amount')
    try:
        result = Decimal(clean(value).replace(',', '').replace('₹', ''))
    except InvalidOperation:
        raise ImportProblem(f'{where}: unreadable target amount') from None
    if not result.is_finite() or not 0 <= result <= Decimal('999999999999.99'):
        raise ImportProblem(f'{where}: target must be finite, nonnegative and fit numeric(14,2)')
    if result != result.quantize(Decimal('0.01')):
        raise ImportProblem(f'{where}: target has more than two decimal places')
    return result


def period(value, fy):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        if value.day != 1:
            raise ImportProblem('Executive month headers must name the first day of the month')
        return value
    text = clean(value)
    if text in MONTHS:
        index = MONTHS.index(text)
        return date(fy + (index >= 9), (index + 3) % 12 + 1, 1)
    for fmt in ('%b-%Y', '%b %Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        if parsed.day != 1:
            raise ImportProblem('Executive month headers must name the first day of the month')
        return parsed
    return None


def source_rows(workbook, fy, assignments):
    if not isinstance(assignments, list) or not assignments:
        raise ImportProblem('Provide explicit person/place/channel/state/city assignments')
    mapping = {}
    for item in assignments:
        if not isinstance(item, dict) or set(item) != ASSIGNMENT_KEYS:
            raise ImportProblem('Each assignment requires exactly person, place, channel, state and city')
        item = {key: clean(value) for key, value in item.items()}
        if any(not item[key] for key in ('person', 'place', 'channel', 'state')):
            raise ImportProblem('Person, source place, channel and state must be explicit')
        if item['channel'] not in pgdata.REGION_OF:
            raise ImportProblem('Assignment contains an unsupported channel')
        key = (item['person'].casefold(), item['place'].casefold())
        if key in mapping:
            raise ImportProblem('Duplicate source assignment')
        mapping[key] = item
    if 'Executive' not in workbook.sheetnames:
        raise ImportProblem('Executive worksheet is missing')
    worksheet = workbook['Executive']
    header = [cell.value for cell in worksheet[1]]
    if len(header) < 3 or clean(header[0]).casefold() != 'person name' or clean(header[1]).casefold() != 'city':
        raise ImportProblem('Executive headers must start with Person Name and City')
    columns = {}
    for index, value in enumerate(header[2:], 2):
        month = period(value, fy)
        if month is None:
            if clean(value):
                raise ImportProblem(f'Executive column {index + 1} has an unrecognized header')
            continue
        if month.year - (month.month < 4) != fy:
            continue
        if month in columns:
            raise ImportProblem('Executive repeats a month in the requested financial year')
        columns[month] = index
    expected = [date(fy + (i >= 9), (i + 3) % 12 + 1, 1) for i in range(12)]
    if set(columns) != set(expected):
        raise ImportProblem('Executive must provide all twelve dated headers for the requested financial year')
    seen, rows, blanks = set(), [], []
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        if len(cells) < 2:
            continue
        key = (clean(cells[0]).casefold(), clean(cells[1]).casefold())
        if key not in mapping:
            continue
        if key in seen:
            raise ImportProblem('Executive repeats a selected person/place row')
        seen.add(key)
        for month in expected:
            index = columns[month]
            coordinate = openpyxl.utils.get_column_letter(index + 1) + str(row_number)
            value = amount(cells[index] if index < len(cells) else None, f'Executive!{coordinate}')
            record = {**mapping[key], 'source_row': row_number, 'source_cell': coordinate,
                      'period_month': month.isoformat()}
            if value is None:
                blanks.append(record)
            else:
                rows.append({**record, 'target_value': format(value, '.2f')})
    if seen != set(mapping):
        missing = [mapping[key]['person'] + ' / ' + mapping[key]['place'] for key in mapping.keys() - seen]
        raise ImportProblem('Selected source assignments not found: ' + ', '.join(sorted(missing)))
    if not rows:
        raise ImportProblem('No entered target amounts in the selected source assignments')
    return rows, blanks


def build_manifest(path, fy, plan_id, source_id, assignments):
    if not 2020 <= fy <= 2098 or plan_id < 1:
        raise ImportProblem('Use an existing plan ID and a supported starting financial year')
    if not re.fullmatch(r'[A-Za-z0-9_-]{20,200}', source_id):
        raise ImportProblem('A valid source spreadsheet ID is required')
    path = Path(path)
    source_bytes = path.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    # Read formulas as formulas: an absent cached result must never be treated
    # as an intentionally blank target. Selected formula cells are rejected.
    workbook = openpyxl.load_workbook(BytesIO(source_bytes), data_only=False)
    try:
        rows, blanks = source_rows(workbook, fy, assignments)
    finally:
        workbook.close()
    manifest = {'source_id': source_id, 'source_sha256': source_hash, 'sheet': 'Executive',
                'fy': fy, 'plan_id': plan_id, 'rows': rows, 'unset_cells': blanks}
    manifest['manifest_sha256'] = hashlib.sha256(json.dumps(manifest, sort_keys=True,
                    ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest()
    return manifest


def resolve_scope(cursor, row):
    resolved = {}
    for dimension, table, column, label in (('person', 'sales_person', 'person_id', 'full_name'),
                                           ('channel', 'channel', 'channel_id', 'code'),
                                           ('state', 'state', 'state_id', 'name')):
        cursor.execute(f'select {column} from master.{table} where lower({label})=lower(%s)', (row[dimension],))
        candidates = cursor.fetchall()
        if len(candidates) != 1:
            raise ImportProblem(f'Unknown or ambiguous {dimension}: {row[dimension]}')
        resolved[dimension + '_id'] = candidates[0][0]
    resolved['city_id'] = None
    if row['city']:
        cursor.execute('select city_id from master.city where lower(name)=lower(%s) and state_id=%s',
                       (row['city'], resolved['state_id']))
        candidates = cursor.fetchall()
        if len(candidates) != 1:
            raise ImportProblem(f"Unknown or ambiguous city in the selected state: {row['city']}")
        resolved['city_id'] = candidates[0][0]
    return resolved


def apply_manifest(connection, manifest, *, apply=False, expected_manifest=None, actor='Executive worksheet import'):
    """Own one transaction; read-only until explicitly called with apply=True."""
    signature = manifest.get('manifest_sha256')
    unsigned = {key: value for key, value in manifest.items() if key != 'manifest_sha256'}
    actual = hashlib.sha256(json.dumps(unsigned, sort_keys=True, ensure_ascii=False,
                      separators=(',', ':')).encode('utf-8')).hexdigest()
    if signature != actual or (apply and expected_manifest != actual):
        raise ImportProblem('Manifest differs from the reviewed preview; no targets were imported')
    report = {'manifest_sha256': actual, 'plan_id': manifest['plan_id'], 'fy': manifest['fy'],
              'inserted': 0, 'unchanged': 0, 'planned_inserts': [], 'unset_cells': manifest['unset_cells']}
    # Writes hold the plan row and target table locks before checking targets;
    # READ COMMITTED sees any editor that finished while those locks waited.
    connection.set_session(readonly=not apply,
                           isolation_level='READ COMMITTED' if apply else 'REPEATABLE READ',
                           autocommit=False)
    try:
        with connection.cursor() as cursor:
            cursor.execute("set local lock_timeout='10s'")
            cursor.execute("set local statement_timeout='30s'")
            cursor.execute('select fy,status from tgt.target_plan where plan_id=%s' + (' for update' if apply else ''),
                           (manifest['plan_id'],))
            plan = cursor.fetchone()
            if plan != (manifest['fy'], 'active'):
                raise ImportProblem('The reviewed plan is no longer active in the requested financial year')
            if apply:
                cursor.execute('lock table tgt.target in share row exclusive mode')
                cursor.execute("select set_config('app.username',%s,true)", (actor,))
            scope_cache, pending, seen = {}, [], set()
            # Blank source cells participate in conflict checks too. They never
            # delete an existing target or silently turn an existing value into zero.
            for row in [*manifest['rows'], *manifest['unset_cells']]:
                scope_key = tuple(row[key] for key in ('person', 'channel', 'state', 'city'))
                if scope_key not in scope_cache:
                    scope_cache[scope_key] = resolve_scope(cursor, row)
                ids = scope_cache[scope_key]
                record_key = (*ids.values(), row['period_month'])
                if record_key in seen:
                    raise ImportProblem('Manifest contains a duplicate resolved target scope')
                seen.add(record_key)
                cursor.execute('''select target_id,channel_id,state_id,city_id,sub_channel_id,district_id,brand_id,
                                          target_value,is_rollup,notes,updated_by
                                   from tgt.target where plan_id=%s and period_month=%s and person_id=%s
                                     and metric='primary_taxable'
                                     and (channel_id is null or channel_id=%s)
                                     and (state_id is null or state_id=%s)
                                     and (city_id is null or %s is null or city_id=%s)''',
                               (manifest['plan_id'], row['period_month'], ids['person_id'], ids['channel_id'],
                                ids['state_id'], ids['city_id'], ids['city_id']))
                existing = cursor.fetchall()
                expected_scope = (ids['channel_id'], ids['state_id'], ids['city_id'], None, None, None)
                for old in existing:
                    same_scope = tuple(old[1:7]) == expected_scope
                    same_value = 'target_value' in row and Decimal(row['target_value']) == old[7]
                    if not same_scope or not same_value or old[8]:
                        raise ImportProblem(f"Existing target conflicts with {row['person']} / {row['place']} / {row['period_month']}; no targets changed")
                if existing:
                    report['unchanged'] += 1
                elif 'target_value' in row:
                    pending.append((row, ids))
                    report['planned_inserts'].append(row)
            if apply:
                for row, ids in pending:
                    notes = (f"Executive worksheet {manifest['source_id']}!Executive!{row['source_cell']}; "
                             f"sha256={manifest['source_sha256']}; place={row['place']}; explicit channel={row['channel']}")
                    cursor.execute('''insert into tgt.target
                        (plan_id,period_month,metric,channel_id,person_id,state_id,city_id,target_value,is_rollup,notes,created_by)
                        values (%s,%s,'primary_taxable',%s,%s,%s,%s,%s,false,%s,%s)''',
                        (manifest['plan_id'], row['period_month'], ids['channel_id'], ids['person_id'],
                         ids['state_id'], ids['city_id'], Decimal(row['target_value']), notes, actor))
                report['inserted'] = len(pending)
                connection.commit()
            else:
                connection.rollback()
        return report
    except Exception:
        connection.rollback()
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--mapping', type=Path, required=True)
    parser.add_argument('--source-id', required=True)
    parser.add_argument('--fy', type=int, required=True)
    parser.add_argument('--plan-id', type=int, required=True)
    parser.add_argument('--report', type=Path, required=True)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--expect-manifest')
    args = parser.parse_args()
    report = None
    try:
        if args.report.resolve() in (args.workbook.resolve(), args.mapping.resolve()):
            raise ImportProblem('Report must not replace the workbook or mapping file')
        # Verify the report destination before any committed database change.
        if args.report.exists() and not args.report.is_file():
            raise ImportProblem('Report destination must be a file')
        with tempfile.TemporaryFile(dir=args.report.parent):
            pass
        manifest = build_manifest(args.workbook, args.fy, args.plan_id, args.source_id,
                                  json.loads(args.mapping.read_text(encoding='utf-8-sig')))
        connection = pgdata.connect()
        try:
            report = apply_manifest(connection, manifest, apply=args.apply, expected_manifest=args.expect_manifest)
        finally:
            connection.close()
        try:
            args.report.write_text(json.dumps({'manifest': manifest, 'result': report}, indent=2,
                                             ensure_ascii=False) + '\n', encoding='utf-8')
        except OSError:
            if args.apply:
                parser.exit(2, f"Targets committed ({report['inserted']} inserted, {report['unchanged']} unchanged), "
                            f"but the report file could not be saved. Manifest {report['manifest_sha256']}.\n")
            raise
        print(json.dumps({'applied': args.apply, 'inserted': report['inserted'],
                          'planned_inserts': len(report['planned_inserts']), 'unchanged': report['unchanged'],
                          'unset_cells': len(report['unset_cells']), 'manifest_sha256': report['manifest_sha256']}))
    except (ImportProblem, ValueError, OSError) as exc:
        parser.exit(1, f'Executive target import rejected: {exc}\n')


if __name__ == '__main__':
    main()
