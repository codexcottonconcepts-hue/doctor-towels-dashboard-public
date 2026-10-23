"""Interactive local PostgreSQL setup. Passwords are entered only in this console.

Creates or upgrades a selected application database, preserves existing role
passwords, verifies a dedicated application login, then saves private config.
"""
import argparse
import getpass
import json
import os
from pathlib import Path
import re
import secrets
import subprocess
import sys
from contextlib import closing
from urllib.parse import quote

import psycopg2
from psycopg2 import sql

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
CONFIG = ROOT / "database.url"
RESULT = ROOT / ".postgres-setup-result.json"


def connection_uri(user, password, database):
    return f"postgresql://{quote(user, safe='')}:{quote(password, safe='')}@127.0.0.1:5432/{quote(database, safe='')}?connect_timeout=5"


def confirm(message):
    return input(message + " [yes/no]: ").strip().lower() == "yes"


def run_script(name, uri, *arguments):
    env = dict(os.environ, DATABASE_URL=uri, PYTHONIOENCODING="utf-8")
    result = subprocess.run([sys.executable, str(HERE / name), *map(str, arguments)], env=env, cwd=ROOT)
    if result.returncode:
        raise RuntimeError("The database step did not complete. The application's saved connection was not changed.")


def protect_file(path):
    """Apply ACLs before writing any password to the new file."""
    env = dict(os.environ, DSTT_CONFIG_PATH=str(path))
    # Windows PowerShell must discover its own modules, not an inherited
    # PowerShell 7 module path from the launching terminal.
    env = {key: value for key, value in env.items() if key.lower() != 'psmodulepath'}
    script = r"""
    $ErrorActionPreference = 'Stop'
    $p = $env:DSTT_CONFIG_PATH
    $acl = Get-Acl -LiteralPath $p
    $acl.SetAccessRuleProtection($true, $false)
    foreach ($rule in @($acl.Access)) { [void]$acl.RemoveAccessRule($rule) }
    $ids = @([Security.Principal.WindowsIdentity]::GetCurrent().User,
             [Security.Principal.SecurityIdentifier]::new('S-1-5-18'),
             [Security.Principal.SecurityIdentifier]::new('S-1-5-32-544'))
    foreach ($id in $ids) {
        if (Test-Path -LiteralPath $p -PathType Container) {
            $rule = [Security.AccessControl.FileSystemAccessRule]::new($id, 'FullControl', 'ContainerInherit, ObjectInherit', 'None', 'Allow')
        } else {
            $rule = [Security.AccessControl.FileSystemAccessRule]::new($id, 'FullControl', 'Allow')
        }
        $acl.AddAccessRule($rule)
    }
    Set-Acl -LiteralPath $p -AclObject $acl
    """
    subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
                   env=env, check=True, capture_output=True)


def save_configuration(uri):
    if CONFIG.exists():
        raise RuntimeError("database.url already exists. It was preserved. Restart the app to use its saved connection, or review it with the administrator.")
    temporary = ROOT / (".database.url." + secrets.token_hex(6) + ".tmp")
    try:
        with temporary.open("x", encoding="utf-8"):
            pass
        protect_file(temporary)
        temporary.write_text(uri, encoding="utf-8")
        # On Windows rename refuses to overwrite a destination that appeared
        # while setup was running. Never replace another setup's credentials.
        temporary.rename(CONFIG)
    finally:
        if temporary.exists():
            temporary.unlink()


def choose_workbook(argument, skip_load):
    if skip_load:
        return None
    value = argument
    if not value:
        print("\nChoose the sales data for the NEW database.")
        print("Enter the path of your approved latest Excel export.")
        print("Type EMPTY only if you want to start a new, empty sales register.")
        downloaded = Path.home() / "Downloads" / "DST & TT DATABASE.xlsx"
        if downloaded.is_file():
            print(f"Type 1 to use: {downloaded}")
        value = input("Excel path, 1, or EMPTY: ").strip().strip('"')
        if value == "1" and downloaded.is_file():
            value = str(downloaded)
    if value.upper() == "EMPTY":
        return None
    workbook = Path(value).expanduser()
    if not value or not workbook.is_file():
        raise RuntimeError("Choose an existing Excel file. No database changes were made.")
    import openpyxl
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        if "Direct Sales Report" not in wb.sheetnames:
            raise RuntimeError("The workbook needs a Direct Sales Report sheet. No database changes were made.")
    finally:
        wb.close()
    return workbook.resolve()


def setup(args):
    if CONFIG.exists():
        raise RuntimeError("A saved database.url already exists. It was preserved. Restart the application to load it before creating another connection.")
    database = args.database or input("Application database name [dstt]: ").strip() or "dstt"
    if not re.fullmatch(r"[a-z][a-z0-9_]{0,48}", database):
        raise RuntimeError("Use a database name starting with a lowercase letter, followed by lowercase letters, digits or underscores.")
    print("\nEnter the PostgreSQL postgres password set during installation.")
    print("It stays on this computer. It is not displayed or saved.")
    password = getpass.getpass("PostgreSQL administrator password: ")
    if not password:
        raise RuntimeError("No password entered. Nothing changed.")
    owner_uri = connection_uri("postgres", password, database)
    try:
        owner = psycopg2.connect(connection_uri("postgres", password, "postgres"))
    except psycopg2.Error:
        raise RuntimeError("PostgreSQL did not accept the connection. Check that the service is running and the password is correct; no database changes were made.") from None
    with closing(owner) as cluster:
        cluster.autocommit = True
        with cluster.cursor() as cur:
            cur.execute("select 1 from pg_database where datname=%s", (database,))
            exists = bool(cur.fetchone())
        schema_exists = False
        if exists:
            with closing(psycopg2.connect(owner_uri)) as conn, conn.cursor() as cur:
                cur.execute("select to_regclass('txn.sales_entry'), to_regclass('tgt.target')")
                tables = cur.fetchone()
                cur.execute("select count(*) from information_schema.tables where table_schema not in ('pg_catalog','information_schema')")
                count = cur.fetchone()[0]
                if count and not all(tables):
                    raise RuntimeError("The selected database contains another or incomplete schema. It was not changed. Choose a new application database name.")
                schema_exists = bool(all(tables))
        workbook = None if schema_exists else choose_workbook(args.excel, args.skip_load)
        if schema_exists:
            print(f"\nExisting application database: {database}. Existing sales and targets will be retained.")
            print("The application's repeatable migrations will be applied. No Excel import will be performed.")
        else:
            print(f"\nA new application schema will be prepared in: {database}")
            print("Sales source: " + (str(workbook) if workbook else "EMPTY register"))
            print("Targets can be added in the Targets page; no old target file is imported automatically.")
        if not confirm("Prepare this database and connect the application"):
            print("Cancelled. No database changes were made.")
            return False
        if not exists:
            with cluster.cursor() as cur:
                cur.execute(sql.SQL("create database {}").format(sql.Identifier(database)))
        run_script("deploy.py", owner_uri, "--upgrade" if schema_exists else "--schema-only")
        if workbook:
            run_script("etl/load_sales.py", owner_uri, workbook, "--dry-run")
            if not confirm("Review the import report above. Import these sales"):
                print("Import cancelled. The database schema remains, but the app is still using its previous source.")
                return False
            run_script("etl/load_sales.py", owner_uri, workbook)
            run_script("deploy.py", owner_uri, "--upgrade")
        login = "dstt_web_" + secrets.token_hex(6)
        app_password = secrets.token_urlsafe(32)
        app_uri = connection_uri(login, app_password, database)
        with cluster.cursor() as cur:
            cur.execute(sql.SQL("create role {} login password %s nosuperuser nocreatedb nocreaterole noreplication inherit in role dstt_app").format(sql.Identifier(login)), (app_password,))
            cur.execute(sql.SQL("grant connect on database {} to {}").format(sql.Identifier(database), sql.Identifier(login)))
        run_script("selftest.py", app_uri)
        with closing(psycopg2.connect(app_uri)) as conn, conn.cursor() as cur:
            cur.execute("select count(*) from rpt.v_sales")
            rows = cur.fetchone()[0]
        save_configuration(app_uri)
        try:
            RESULT.write_text(json.dumps({"ok": True, "database": database, "sales_rows": rows}), encoding="utf-8")
        except OSError:
            print("The connection is saved. The optional setup status file could not be written.")
        print(f"\nConnection saved securely. PostgreSQL contains {rows:,} sales lines.")
        print("Restart the dashboard app, then refresh Sales entry in your browser.")
        print("Existing database passwords, firewall settings and authentication rules were not changed.")
        return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database")
    parser.add_argument("--excel")
    parser.add_argument("--skip-load", action="store_true")
    args = parser.parse_args()
    print("Doctor Towels - connect local PostgreSQL\n")
    try:
        return 0 if setup(args) else 1
    except (KeyboardInterrupt, EOFError):
        print("\nSetup cancelled. The saved application connection was not changed.")
        return 1
    except RuntimeError as exc:
        print("\n" + str(exc))
        return 1
    except Exception:
        print("\nSetup did not finish. The saved connection was not changed. Ask for help with the failed step; do not share database passwords.")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
