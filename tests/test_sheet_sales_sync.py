"""Focused normalization, identity and three-way merge regression checks."""
from datetime import date
from decimal import Decimal
import os
from pathlib import Path
import unittest
from unittest.mock import MagicMock, patch
import uuid

import openpyxl
import psycopg2
from psycopg2.extensions import parse_dsn

from sheet_sales_sync import (HEADERS, OPTIONAL_HEADERS, SyncValidationError,
    Masters, canonical, merge_payload, number, parse_workbook, sync_workbook)


def workbook(rows=None, extra=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Direct Sales Report'
    fields = [k for k in HEADERS if k not in OPTIONAL_HEADERS] + list(extra)
    header = [HEADERS[k][0] for k in fields]
    header[fields.index('brand')] = 'Company'
    header[fields.index('sub_channel')] = 'Sub Channels'
    ws.append(header)
    base = dict(source_key=1, invoice_date=date(2026, 4, 4), invoice_no='SYNC-TEST',
                customer='Sync test customer', channel='D&R - S', person='Sync test person',
                city='Chennai', state='Tamil Nadu', sub_channel='Direct', brand='Doctor Towel',
                order_type='Re-Order', primary_qty=10, primary_taxable=100, primary_with_tax=118)
    for row in rows if rows is not None else [{}]:
        values = {**base, **row}
        ws.append([values.get(k) for k in fields])
    return wb


def serial_only_row(key):
    return {**dict.fromkeys(HEADERS), 'source_key': key}


class SheetParsingTests(unittest.TestCase):
    def test_current_headers_and_stable_serial(self):
        parsed = parse_workbook(workbook([{'source_key': 1.0}]))[0]
        self.assertEqual(parsed['source_key'], '1')
        self.assertEqual(parsed['channel'], 'D&R-S')
        self.assertEqual(parsed['order_type'], 'Re-order')
        self.assertNotIn('courier', parsed)

    def test_duplicate_or_missing_identity_rejects_whole_source(self):
        for rows in ([{'source_key': 1}, {'source_key': '1.0'}], [{'source_key': None}], [{'source_key': 'customer-secret'}]):
            with self.subTest(rows=rows), self.assertRaises(SyncValidationError) as raised:
                parse_workbook(workbook(rows))
            self.assertNotIn('customer-secret', str(raised.exception))

    def test_serial_only_rows_do_not_reserve_a_real_invoice_serial(self):
        for rows in ([serial_only_row(1), {'source_key': 1}, serial_only_row('1.0')],
                     [{'source_key': 1}, serial_only_row(1), serial_only_row('1.0')]):
            with self.subTest(rows=rows):
                diagnostics = {}
                parsed = parse_workbook(workbook(rows), diagnostics=diagnostics)
                self.assertEqual([row['source_key'] for row in parsed], ['1'])
                self.assertEqual(diagnostics['source_rows'], 3)
                self.assertEqual(diagnostics['invoice_rows'], 1)
                self.assertEqual(diagnostics['ignored_blank_rows'], 2)
                self.assertEqual([row['source_key'] for row in diagnostics['ignored_rows']], ['1', '1'])

    def test_real_duplicates_report_both_sheet_rows_and_keep_blank_diagnostics(self):
        diagnostics = {}
        with self.assertRaises(SyncValidationError) as raised:
            parse_workbook(workbook([{'source_key': 999}, serial_only_row(987),
                                     {'source_key': '999.0', 'invoice_no': 'ANOTHER-INVOICE'}]),
                           diagnostics=diagnostics)
        self.assertEqual(raised.exception.issues, [{'source_key': '999', 'code': 'duplicate_source_key',
                                                  'sheet_row': 4, 'sheet_rows': [2, 4]}])
        self.assertIn('worksheet rows 2, 4', raised.exception.public_message)
        self.assertNotIn('ANOTHER-INVOICE', raised.exception.public_message)
        self.assertEqual(raised.exception.diagnostics, diagnostics)
        self.assertEqual(diagnostics['source_rows'], 3)
        self.assertEqual(diagnostics['invoice_rows'], 2)
        self.assertEqual(diagnostics['ignored_blank_rows'], 1)
        diagnostics['ignored_rows'][0]['source_key'] = 'changed'
        self.assertEqual(raised.exception.diagnostics['ignored_rows'][0]['source_key'], '987')

    def test_missing_invoice_with_any_other_data_is_still_invalid(self):
        for row in ({'invoice_no': None}, {**serial_only_row(1), 'primary_qty': 0},
                    {**serial_only_row(1), 'remarks': 'Needs invoice'}):
            with self.subTest(row=row), self.assertRaises(SyncValidationError) as raised:
                parse_workbook(workbook([row], extra=('remarks',)))
            self.assertEqual(raised.exception.issues[0]['code'], 'missing_invoice_no')
            self.assertEqual(raised.exception.diagnostics['ignored_blank_rows'], 0)

    def test_unmapped_column_content_cannot_be_ignored_as_a_placeholder(self):
        wb = workbook([serial_only_row(1)])
        column = wb.active.max_column + 1
        wb.active.cell(1, column).value = 'Other source information'
        wb.active.cell(2, column).value = 'Data still needs review'
        with self.assertRaises(SyncValidationError) as raised:
            parse_workbook(wb)
        self.assertEqual(raised.exception.issues[0]['code'], 'missing_invoice_no')
        self.assertEqual(raised.exception.diagnostics['ignored_blank_rows'], 0)

    def test_unmapped_month_marker_without_record_fields_is_ignored(self):
        # The real source begins with Month-only marker rows. They have no
        # serial or mapped record fields and are not invoice placeholders.
        wb = workbook([serial_only_row(None), {'source_key': 1}])
        column = wb.active.max_column + 1
        wb.active.cell(1, column).value = 'Month'
        wb.active.cell(2, column).value = date(2026, 4, 1)
        diagnostics = {}
        parsed = parse_workbook(wb, diagnostics=diagnostics)
        self.assertEqual([row['source_key'] for row in parsed], ['1'])
        self.assertEqual(diagnostics['source_rows'], 1)
        self.assertEqual(diagnostics['invoice_rows'], 1)
        self.assertEqual(diagnostics['ignored_blank_rows'], 0)
        self.assertEqual(diagnostics['ignored_rows'], [])

    def test_invalid_serial_only_rows_are_not_silently_ignored(self):
        for key in (0, -1, 'NaN', 'customer-secret'):
            with self.subTest(key=key), self.assertRaises(SyncValidationError) as raised:
                parse_workbook(workbook([serial_only_row(key)]))
            self.assertEqual(raised.exception.issues[0]['code'], 'missing_or_invalid_source_key')
            self.assertEqual(raised.exception.diagnostics['ignored_blank_rows'], 0)
            self.assertNotIn('customer-secret', str(raised.exception))

    def test_all_placeholders_are_rejected_as_empty_with_diagnostics(self):
        with self.assertRaises(SyncValidationError) as raised:
            parse_workbook(workbook([serial_only_row(987), serial_only_row(987)]))
        self.assertEqual(raised.exception.issues, [{'code': 'empty_source'}])
        self.assertEqual(raised.exception.diagnostics['source_rows'], 2)
        self.assertEqual(raised.exception.diagnostics['invoice_rows'], 0)
        self.assertEqual(raised.exception.diagnostics['ignored_blank_rows'], 2)

    def test_duplicate_real_rows_fail_before_database_or_master_work(self):
        conn = MagicMock()
        with patch('sheet_sales_sync.Masters') as masters, self.assertRaises(SyncValidationError):
            sync_workbook(workbook([{'source_key': 1}, serial_only_row(1), {'source_key': '1.0'}]),
                          'test-source', conn)
        conn.cursor.assert_not_called()
        conn.__enter__.assert_not_called()
        masters.assert_not_called()

    def test_validation_message_indicates_additional_issues(self):
        with self.assertRaises(SyncValidationError) as raised:
            parse_workbook(workbook([{'source_key': key, 'invoice_no': None} for key in range(1, 13)]))
        self.assertEqual(len(raised.exception.issues), 12)
        self.assertTrue(raised.exception.public_message.endswith('; 2 more issue(s)'))

    def test_invalid_numeric_never_becomes_zero(self):
        for value in ('typo', 'NaN', 'Infinity', '-1', True, '1000000000000'):
            with self.subTest(value=value), self.assertRaises(SyncValidationError):
                parse_workbook(workbook([{'primary_taxable': value}]))

    def test_formula_money_precision_matches_postgresql(self):
        self.assertEqual(number(100.125), Decimal('100.13'))
        self.assertEqual(number(None), Decimal('0.00'))
        self.assertEqual(number('1,234.50'), Decimal('1234.50'))

    def test_truncated_or_empty_workbook_rejected(self):
        wb = workbook()
        wb.active.cell(1, 4).value = 'unexpected'
        with self.assertRaises(SyncValidationError):
            parse_workbook(wb)
        with self.assertRaises(SyncValidationError):
            parse_workbook(workbook([]))

    def test_optional_column_zero_padded_docket(self):
        parsed = parse_workbook(workbook([{'courier': 'dtdc', 'courier_docket_no': '001234'}], extra=('courier', 'courier_docket_no')))[0]
        self.assertEqual(parsed['courier'], 'DTDC')
        self.assertEqual(parsed['courier_docket_no'], '001234')

    def test_invalid_dates_and_credit_are_rejected(self):
        for changes in ({'invoice_date': '2026-02-30'}, {'stock_supplied_on': 'typo'},
                        {'credit_days': '1.2'}, {'credit_days': 'NaN'},
                        {'stock_supplied_on': date(2026, 4, 5), 'payment_received_on': date(2026, 4, 4)}):
            with self.subTest(changes=changes), self.assertRaises(SyncValidationError):
                parse_workbook(workbook([changes]))


class MergeTests(unittest.TestCase):
    def test_unchanged_sheet_preserves_local_edits(self):
        self.assertEqual(merge_payload({'primary_taxable': '100'}, {'primary_taxable': 100},
                                      {'primary_taxable': Decimal('101.00')}), ({}, []))

    def test_independent_edits_merge_and_same_field_conflicts(self):
        baseline = {'primary_taxable': '100', 'pi_number': 'A'}
        source = {'primary_taxable': '120', 'pi_number': 'A'}
        self.assertEqual(merge_payload(baseline, source, {'primary_taxable': 100, 'pi_number': 'B'}),
                         ({'primary_taxable': '120.00'}, []))
        self.assertEqual(merge_payload(baseline, source, {'primary_taxable': 110, 'pi_number': 'A'})[1],
                         ['primary_taxable'])

    def test_added_dispatch_column_cannot_erase_dashboard_values(self):
        for source in ({'courier_id': None}, {'courier_id': 2}):
            with self.subTest(source=source):
                self.assertEqual(merge_payload({}, source, {'courier_id': 3})[1], ['courier_id'])
        self.assertEqual(merge_payload({}, {'courier_id': 2}, {'courier_id': None}), ({'courier_id': 2}, []))

    def test_coupled_geography_edits_conflict(self):
        baseline = {'city_id': 1, 'state_id': 1, 'district_id': 1}
        source = {**baseline, 'state_id': 2}
        current = {**baseline, 'city_id': 3}
        self.assertEqual(merge_payload(baseline, source, current)[1], ['state_id'])

    def test_audit_json_numeric_matches_ledger_decimal(self):
        self.assertEqual(canonical({'primary_taxable': 100}), canonical({'primary_taxable': Decimal('100.00')}))


class SyncPlaceholderProtectionTests(unittest.TestCase):
    def test_read_only_preview_avoids_row_locks_and_writes(self):
        payload = {'invoice_no': 'UNCHANGED', 'brand_id': 1}
        for dry_run in (True, False):
            with self.subTest(dry_run=dry_run):
                conn = MagicMock()
                master_cursor, row_cursor = MagicMock(), MagicMock()
                conn.cursor.side_effect = [master_cursor, row_cursor]
                master_cursor.__enter__.return_value.fetchone.return_value = (True,)
                cur = row_cursor.__enter__.return_value
                cur.fetchall.side_effect = [
                    [{'source_id': 'test-source', 'source_key': '1', 'entry_id': 1,
                      'source_payload': payload.copy()}],
                    [{**payload, 'entry_id': 1, 'is_deleted': False, 'source': 'sheets_migration'}]]
                with patch('sheet_sales_sync.Masters.payload', return_value=payload.copy()):
                    result = sync_workbook(workbook(), 'test-source', conn, dry_run=dry_run)
                queries = [call.args[0].lower() for call in cur.execute.call_args_list]
                self.assertEqual(any('for update' in query for query in queries), not dry_run)
                if dry_run:
                    self.assertFalse(any(query.startswith(('insert', 'update', 'delete')) for query in queries))
                self.assertEqual(result['unchanged'], 1)

    def test_linked_placeholder_retains_ledger_and_baseline_and_is_not_missing(self):
        payload = {'invoice_no': 'UNCHANGED', 'brand_id': 1}
        links = [{'source_id': 'test-source', 'source_key': str(key), 'entry_id': key,
                  'source_payload': payload.copy()} for key in (1, 2)]
        entries = [{**payload, 'entry_id': key, 'is_deleted': False, 'source': 'sheets_migration'}
                   for key in (1, 2)]
        conn = MagicMock()
        master_cursor, row_cursor = MagicMock(), MagicMock()
        conn.cursor.side_effect = [master_cursor, row_cursor]
        master_cursor.__enter__.return_value.fetchone.return_value = (True,)
        cur = row_cursor.__enter__.return_value
        cur.fetchall.side_effect = [links, entries]
        with patch('sheet_sales_sync.Masters.payload', return_value=payload.copy()):
            result = sync_workbook(workbook([serial_only_row(1), {'source_key': 2}]),
                                   'test-source', conn)
        self.assertEqual(result['source_rows'], 1)
        self.assertEqual(result['source_total_rows'], 2)
        self.assertEqual(result['invoice_rows'], 1)
        self.assertEqual(result['ignored_blank_rows'], 1)
        self.assertEqual(result['conflicts'], 1)
        self.assertEqual(result['missing_source'], 0)
        self.assertEqual(result['updated'], 0)
        self.assertEqual(result['inserted'], 0)
        self.assertIn({'source_key': '1', 'code': 'linked_source_row_blank', 'sheet_rows': [2]}, result['issues'])
        writes = [call.args for call in cur.execute.call_args_list
                  if call.args[0].lower().startswith(('insert', 'update', 'delete'))]
        self.assertEqual(len(writes), 1)
        self.assertTrue(writes[0][0].startswith('insert into aux.sheet_sales_row'))
        self.assertEqual(writes[0][1][1], '2')
        self.assertEqual(links[0]['source_payload'], payload)

    def test_placeholder_relinks_unique_invoice_without_inserting_it_twice(self):
        payload = {'invoice_no': 'EXISTING-INVOICE', 'brand_id': 1}
        conn = MagicMock()
        master_cursor, row_cursor = MagicMock(), MagicMock()
        conn.cursor.side_effect = [master_cursor, row_cursor]
        master_cursor.__enter__.return_value.fetchone.return_value = (True,)
        cur = row_cursor.__enter__.return_value
        cur.rowcount = 1
        cur.fetchall.side_effect = [
            [{'source_id': 'test-source', 'source_key': '1', 'entry_id': 1, 'source_payload': payload.copy()}],
            [{**payload, 'entry_id': 1, 'is_deleted': False, 'source': 'sheets_migration'}]]
        with patch('sheet_sales_sync.Masters.payload', return_value=payload.copy()):
            result = sync_workbook(workbook([serial_only_row(1), {'source_key': 2}]), 'test-source', conn)
        self.assertEqual(result['relinked'], 1)
        self.assertEqual(result['ignored_blank_rows'], 1)
        self.assertEqual(result['inserted'], 0)
        self.assertEqual(result['updated'], 0)
        self.assertEqual(result['conflicts'], 0)
        self.assertEqual(result['missing_source'], 0)
        writes = [call.args for call in cur.execute.call_args_list
                  if call.args[0].lower().startswith(('insert', 'update', 'delete'))]
        self.assertEqual(len(writes), 3)
        self.assertFalse(any(call[0].startswith(('insert into txn.', 'update txn.', 'delete')) for call in writes))
        self.assertEqual(writes[-1][1][1], '2')


@unittest.skipUnless(os.environ.get('SHEET_SYNC_TEST_DATABASE_URL_FILE'), 'Requires a dedicated disposable synchronization database')
class PostgresSyncTests(unittest.TestCase):
    """One scenario owns its disposable database; never uses DATABASE_URL."""
    def setUp(self):
        dsn = Path(os.environ['SHEET_SYNC_TEST_DATABASE_URL_FILE']).read_text().strip()
        self.assertIn('_test', parse_dsn(dsn)['dbname'])
        self.conn = psycopg2.connect(dsn)
        self.addCleanup(self.conn.close)

    def seed(self, row, source='sheets_migration', line=1):
        with self.conn, self.conn.cursor() as cur:
            payload = Masters(cur, False).payload(parse_workbook(workbook([row]))[0])
            payload.update(source=source, line_no=line)
            cur.execute(f"insert into txn.sales_entry({','.join(payload)}) values ({','.join(['%s']*len(payload))}) returning entry_id", list(payload.values()))
            return cur.fetchone()[0]

    def test_audit_bootstrap_and_duplicate_invoice_reordering(self):
        invoice = 'SYNC-BOOT-' + uuid.uuid4().hex
        first = {'invoice_no': invoice, 'invoice_date': date(2026, 1, 4)}
        second = {**first, 'source_key': 2, 'primary_qty': 20}
        entry1 = self.seed(first)
        entry2 = self.seed(second, line=2)
        with self.conn, self.conn.cursor() as cur:
            cur.execute('update txn.sales_entry set invoice_date=%s where entry_id=%s', (date(2026, 4, 4), entry1))
        source = 'bootstrap-' + uuid.uuid4().hex
        result = sync_workbook(workbook([second, first]), source, self.conn, dry_run=True)
        self.assertEqual(result['attached'], 2)
        self.assertEqual(result['bootstrap_preserved'], 1)
        result = sync_workbook(workbook([second, first]), source, self.conn)
        self.assertEqual(result['inserted'], 0)
        with self.conn, self.conn.cursor() as cur:
            cur.execute('select source_key,entry_id from aux.sheet_sales_row where source_id=%s', (source,))
            self.assertEqual(dict(cur.fetchall()), {'1': entry1, '2': entry2})
            cur.execute('select invoice_date from txn.sales_entry where entry_id=%s', (entry1,))
            self.assertEqual(cur.fetchone()[0], date(2026, 4, 4))
        self.assertEqual(sync_workbook(workbook([first, second]), source, self.conn)['preserved_local'], 1)

    def test_web_rows_and_deleted_legacy_are_protected(self):
        for source_kind in ('web', 'sheets_migration'):
            row = {'invoice_no': 'SYNC-PROTECT-' + uuid.uuid4().hex}
            entry = self.seed(row, source=source_kind)
            if source_kind == 'sheets_migration':
                with self.conn, self.conn.cursor() as cur:
                    cur.execute('update txn.sales_entry set is_deleted=true where entry_id=%s', (entry,))
            with self.assertRaises(SyncValidationError):
                sync_workbook(workbook([row]), 'protect-' + uuid.uuid4().hex, self.conn)
            with self.conn, self.conn.cursor() as cur:
                cur.execute('select count(*) from txn.sales_entry where invoice_no=%s', (row['invoice_no'],))
                self.assertEqual(cur.fetchone()[0], 1)

    def test_mapped_soft_deleted_row_stays_deleted(self):
        row = {'invoice_no': 'SYNC-DELETE-' + uuid.uuid4().hex}
        source = 'delete-' + uuid.uuid4().hex
        sync_workbook(workbook([row]), source, self.conn)
        with self.conn, self.conn.cursor() as cur:
            cur.execute('update txn.sales_entry set is_deleted=true where entry_id=(select entry_id from aux.sheet_sales_row where source_id=%s)', (source,))
        result = sync_workbook(workbook([row]), source, self.conn)
        self.assertEqual(result['conflicts'], 1)
        self.assertEqual(result['issues'][0]['code'], 'database_row_deleted')
        self.assertEqual(result['inserted'], 0)

    def test_linked_placeholder_preserves_entry_and_accepted_source(self):
        source = 'blank-protect-' + uuid.uuid4().hex
        first = {'source_key': 1, 'invoice_no': 'SYNC-BLANK-' + uuid.uuid4().hex}
        second = {'source_key': 2, 'invoice_no': 'SYNC-KEEP-' + uuid.uuid4().hex}
        sync_workbook(workbook([first, second]), source, self.conn)
        query = ('select row_to_json(s), l.source_payload, l.source_fingerprint, l.checked_at '
                 'from aux.sheet_sales_row l join txn.sales_entry s using(entry_id) '
                 'where l.source_id=%s and l.source_key=%s')
        with self.conn, self.conn.cursor() as cur:
            cur.execute(query, (source, '1'))
            before = cur.fetchone()
        result = sync_workbook(workbook([serial_only_row(1), second]), source, self.conn)
        self.assertEqual(result['ignored_blank_rows'], 1)
        self.assertEqual(result['conflicts'], 1)
        self.assertEqual(result['missing_source'], 0)
        self.assertEqual(result['updated'], 0)
        self.assertEqual(result['inserted'], 0)
        with self.conn, self.conn.cursor() as cur:
            cur.execute(query, (source, '1'))
            self.assertEqual(cur.fetchone(), before)

    def test_independent_invoice_brand_edits_allocate_merged_destination(self):
        for local_field in ('invoice_no', 'brand_id'):
            with self.subTest(local_field=local_field):
                original_invoice = 'SYNC-FROM-' + uuid.uuid4().hex
                destination_invoice = 'SYNC-TO-' + uuid.uuid4().hex
                source = 'merge-identity-' + uuid.uuid4().hex
                original = {'invoice_no': original_invoice}
                sync_workbook(workbook([original]), source, self.conn)
                # Destination already contains a web-entered TT line. The
                # merged sheet row must receive its next available line number.
                self.seed({'invoice_no': destination_invoice, 'brand': 'Thirsty Towel'}, source='web')
                with self.conn, self.conn.cursor() as cur:
                    cur.execute("select brand_id from master.brand where name='Thirsty Towel'")
                    brand = cur.fetchone()[0]
                    cur.execute('select entry_id from aux.sheet_sales_row where source_id=%s', (source,))
                    entry_id = cur.fetchone()[0]
                    if local_field == 'invoice_no':
                        cur.execute('update txn.sales_entry set invoice_no=%s where entry_id=%s', (destination_invoice, entry_id))
                        incoming = {**original, 'brand': 'Thirsty Towel'}
                    else:
                        cur.execute('update txn.sales_entry set brand_id=%s where entry_id=%s', (brand, entry_id))
                        incoming = {**original, 'invoice_no': destination_invoice}
                self.assertEqual(sync_workbook(workbook([incoming]), source, self.conn)['updated'], 1)
                with self.conn, self.conn.cursor() as cur:
                    cur.execute('select invoice_no,brand_id,line_no from txn.sales_entry where entry_id=%s', (entry_id,))
                    self.assertEqual(cur.fetchone(), (destination_invoice, brand, 2))

    def test_end_to_end(self):
        conn = self.conn
        source = 'test-' + uuid.uuid4().hex
        invoice = 'SYNC-' + uuid.uuid4().hex
        first = {'invoice_no': invoice}
        preview = sync_workbook(workbook([first]), source, conn, dry_run=True)
        self.assertEqual(preview['inserted'], 1)
        with conn, conn.cursor() as cur:
            cur.execute('select count(*) from aux.sheet_sales_row where source_id=%s', (source,))
            self.assertEqual(cur.fetchone()[0], 0)
        self.assertEqual(sync_workbook(workbook([first]), source, conn)['inserted'], 1)
        self.assertEqual(sync_workbook(workbook([first]), source, conn)['unchanged'], 1)
        # A second SKU on the same invoice is a distinct permanent S.NO.
        second = {**first, 'source_key': 2, 'primary_qty': 20}
        self.assertEqual(sync_workbook(workbook([second, first]), source, conn)['inserted'], 1)
        with conn, conn.cursor() as cur:
            cur.execute('select source_key,entry_id from aux.sheet_sales_row where source_id=%s', (source,))
            ids = dict(cur.fetchall())
            cur.execute('update txn.sales_entry set courier_docket_no=%s,primary_qty=11 where entry_id=%s', ('00123', ids['1']))
        unchanged = sync_workbook(workbook([second, first]), source, conn)
        self.assertEqual(unchanged['preserved_local'], 1)
        changed = {**first, 'primary_taxable': 105}
        self.assertEqual(sync_workbook(workbook([changed, second]), source, conn)['updated'], 1)
        with conn, conn.cursor() as cur:
            cur.execute('select primary_qty,courier_docket_no,primary_taxable from txn.sales_entry where entry_id=%s', (ids['1'],))
            self.assertEqual(cur.fetchone(), (Decimal('11'), '00123', Decimal('105')))
        conflict = {**changed, 'primary_qty': 12}
        self.assertEqual(sync_workbook(workbook([conflict, second]), source, conn)['conflicts'], 1)
        self.assertEqual(sync_workbook(workbook([changed]), source, conn)['missing_source'], 1)
        with self.assertRaises(SyncValidationError):
            sync_workbook(workbook([changed, {**second, 'primary_taxable': 'bad'}]), source, conn)
        with conn, conn.cursor() as cur:
            cur.execute('select count(*) from txn.sales_entry where entry_id=any(%s) and not is_deleted', (list(ids.values()),))
            self.assertEqual(cur.fetchone()[0], 2)
