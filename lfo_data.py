"""LFO branch sales and separate secondary targets from LFO Distributed.

This module does not download sources or touch the primary invoice ledger.
Callers supply a data_only openpyxl workbook and a PostgreSQL connection.
import_workbook and save_monthly_sale own their transaction; fetch_dashboard is
read-only and can participate in the caller's dashboard snapshot transaction.
"""
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_FLOOR
import hashlib
import json
import os
import re
import unicodedata
import xml.etree.ElementTree as ET

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.datetime import from_excel

from psycopg2.extras import Json

SHEET = 'LFO Distributed'
SOURCE_ID = os.environ.get('SALES_SOURCE_ID', '').strip()
MONTHS = ('Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec','Jan','Feb','Mar')
MONTH_NUMBERS = (4,5,6,7,8,9,10,11,12,1,2,3)
GROUPS = (('LULU','LULU','lulu'), ('NILKAMAL','Nilkamal Homes','nilkamal'),
          ('RELIANCE','RELIANCE','reliance'))
TARGET_BASIS = 'distributor_monthly_sheet'
STORE_TARGET_BASIS = 'sheet_branch_plan'
# Add only identities verified against your own workbook. Private branch
# mappings are intentionally absent from this public code copy.
BRANCH_TARGET_ALIASES = {}
SALES_BASIS_NOTE = ('Nilkamal amounts are marked gross sales; the sheet does not say whether '
                    'all distributors include GST. Amounts are shown as entered.')
TOTALS_NOTE = "Sales totals use each store's monthly entries because the sheet's summary totals differ."

# L-LFO keeps one flat table -- Group, Location, targets, then one column per
# month -- where LFO Distributed keeps a section per distributor. Same tables,
# separated by lfo.distributor.channel (migration 14).
LLFO_SHEET = 'L-LFO Distributed'
LLFO_GID = int(os.environ.get('LLFO_SHEET_GID', '0') or '0')
LFO_GID = int(os.environ.get('LFO_SHEET_GID', '0') or '0')
LLFO_TARGET_BASIS = 'sum_of_branch_monthly_targets'
LLFO_STORE_TARGET_BASIS = 'sheet_branch_monthly_target'
LLFO_MONTH_WORDS = {'apr':4,'may':5,'jun':6,'jul':7,'aug':8,'sep':9,'oct':10,'nov':11,'dec':12,
                    'jan':1,'feb':2,'mar':3}
LLFO_NOT_IMPORTED_NOTE = ('Primary Target, Primary Achieved, the single Secondary Achieved column and '
                          'Ach % are not imported. Secondary sales come only from the monthly columns; '
                          'primary sales and targets come from the invoice ledger and the primary plan.')
CHANNELS = {
    'LFO': {'sheet':SHEET, 'gid':LFO_GID, 'targetBasis':TARGET_BASIS, 'storeTargetBasis':STORE_TARGET_BASIS,
            'description':'Distributor Monthly Target columns; yearly targets are the sum of April to March.',
            'storeSplit':'Exact Month and Year amounts from the branch target tables; no scaling or equal split.'},
    'L-LFO': {'sheet':LLFO_SHEET, 'gid':LLFO_GID, 'targetBasis':LLFO_TARGET_BASIS,
              'storeTargetBasis':LLFO_STORE_TARGET_BASIS,
              'description':"Each group's monthly target is the sum of its branches' Secondary Target; "
                            'yearly targets are the sum of April to March.',
              'storeSplit':"Each branch's own Secondary Target per month; its year is twelve times that "
                           'amount because the sheet carries no Year cell.'},
}
# Targets saved in the Target planner carry this source. A sheet import never
# changes a group and financial year that has any, so a planned year stays as
# planned even after the Distributed tab moves on to that year.
PLANNER_SOURCE = 'dashboard'
PLANNER_SHEET = 'Target planner'
PLANNER_BASIS = 'target_planner'
# These channels have no Distributed tab: their sub-channels, branches and
# targets are set up in the Target planner (migration 15 allows them).
for _code in ('D&R-S', 'D&R-N', 'Saloon'):
    CHANNELS[_code] = {'sheet':None, 'gid':None, 'targetBasis':PLANNER_BASIS, 'storeTargetBasis':PLANNER_BASIS,
                       'description':'Sub-channel monthly targets set in the Target planner; yearly targets are the sum of April to March.',
                       'storeSplit':"Each branch's monthly target is set in the Target planner; its year is twelve times that amount."}
del _code
PLANNED_NOTE = ('{name}: FY {fy}–{yy} targets are set in the Target planner, so this sheet’s '
                'targets for that year were not imported. Its sales were imported as usual.')


class LFOValidationError(ValueError):
    def __init__(self, message):
        self.public_message = str(message)
        super().__init__(self.public_message)


class LFOConflictError(ValueError):
    public_message = 'This monthly sale changed after you opened it. Reload it before saving.'
    def __init__(self):
        super().__init__(self.public_message)


def normalized_name(value):
    return ' '.join(str(value or '').split()).casefold()


def money(value, nullable=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        if nullable:
            return None
        raise LFOValidationError('Enter a monthly sales amount; use 0 for no sales.')
    if isinstance(value, bool):
        raise LFOValidationError('The amount must be a non-negative number.')
    try:
        result = Decimal(str(value).replace(',', '').replace('\u20b9','').strip())
        if not result.is_finite() or result < 0 or result >= Decimal('1000000000000'):
            raise InvalidOperation
        result = result.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        if result >= Decimal('1000000000000'):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise LFOValidationError('The amount must be a finite non-negative number below 1 trillion.') from None
    return result


def period(value):
    if isinstance(value, datetime):
        value = value.date()
    elif isinstance(value, str):
        try:
            value = date.fromisoformat(value)
        except ValueError:
            raise LFOValidationError('Select a valid sales month.') from None
    if not isinstance(value, date) or value.day != 1 or not 2000 <= value.year <= 2100:
        raise LFOValidationError('Select a valid sales month (first day of the month).')
    return value


def _fy(value):
    match = re.search(r'(?<!\d)(20\d{2}|\d{2})\s*[-–/]\s*(20\d{2}|\d{2})(?!\d)', str(value or ''))
    if not match:
        return None
    first, last = (int(x) for x in match.groups())
    first = first + 2000 if first < 100 else first
    last = last + 2000 if last < 100 else last
    if last != first + 1:
        raise LFOValidationError('The LFO sheet financial-year heading is inconsistent.')
    return first


def _source_month(value):
    if isinstance(value, (date, datetime)):
        return value.month, value.year
    txt = str(value or '').strip()
    for fmt in ('%b-%Y','%B-%Y','%b-%y','%B-%y','%Y-%m-%d','%d-%m-%Y','%d/%m/%Y'):
        try:
            parsed = datetime.strptime(txt, fmt)
            return parsed.month, parsed.year
        except ValueError:
            pass
    for index, name in enumerate(MONTHS):
        if txt.casefold() == name.casefold():
            return MONTH_NUMBERS[index], None
    return None, None


def _month_date(value, fy, expected, warnings):
    month, source_year = _source_month(value)
    if month != expected:
        raise LFOValidationError('LFO monthly columns must run April through March without duplicates.')
    result = date(fy if month >= 4 else fy+1, month, 1)
    if source_year is not None and source_year != result.year:
        if month <= 3 and source_year == fy:
            warning = f'January to March dates use {fy+1} to match financial year {fy}–{str(fy+1)[-2:]}.'
            if warning not in warnings:
                warnings.append(warning)
        else:
            raise LFOValidationError('A monthly date does not match its financial-year section.')
    return result


class _MonetaryCells:
    """Recover numeric XLSX amounts mistakenly displayed with date formatting.

    This is confined to monetary coordinates after the source layout has been
    validated. An original numeric XML value is required; an unverified Python
    date, an ISO date cell, or a date-looking string is never an amount. Reading
    the raw decimal also avoids Excel epoch/leap-day and float round-trip loss.
    """
    def __init__(self,workbook,worksheet):
        self.workbook,self.worksheet = workbook,worksheet
        self.raw = None
        self.recovered = []

    def __call__(self,cell,nullable=True):
        value = cell.value if cell is not None else None
        if not isinstance(value,(date,datetime,time,timedelta)):
            return money(value,nullable=nullable)
        address = getattr(cell,'coordinate','unknown')
        message = f'{address} contains a date in a sales or target amount. Restore its numeric amount and use a number format.'
        if cell.data_type!='d' or not is_date_format(cell.number_format):
            raise LFOValidationError(message)
        archive = getattr(self.workbook,'_archive',None)
        path = getattr(self.worksheet,'_worksheet_path',None)
        if archive is None or path is None:
            raise LFOValidationError(message)
        if self.raw is None:
            ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            root = ET.fromstring(archive.read(path))
            self.raw = {}
            for node in root.iter(ns+'c'):
                raw_value = node.find(ns+'v')
                if node.get('t','n')=='n' and raw_value is not None:
                    self.raw[node.get('r')] = raw_value.text
        raw = self.raw.get(address)
        if raw is None:
            raise LFOValidationError(message)
        amount = money(raw)
        try:
            interpreted = from_excel(float(Decimal(raw)),self.workbook.epoch,
                                     timedelta=isinstance(value,timedelta))
        except (ValueError,OverflowError,InvalidOperation):
            raise LFOValidationError(message) from None
        if interpreted!=value:
            raise LFOValidationError(message)
        self.recovered.append(address)
        return amount


def parse_workbook(workbook):
    """Validate source structure and return distributors, stores, weights, sales.

    Store actuals come only from the twelve monthly cells in each branch row.
    Blank monthly cells are retained as None and never turned into zero. Totals,
    SOH, source store-count plans and yearly target text are never sales inputs.
    """
    if SHEET not in workbook.sheetnames:
        raise LFOValidationError('The workbook has no LFO Distributed worksheet.')
    worksheet = workbook[SHEET]
    rows = list(worksheet.iter_rows())
    def source_cell(r,c):
        return rows[r-1][c-1] if 0 < r <= len(rows) and c <= len(rows[r-1]) else None
    def cell(r,c):
        entry = source_cell(r,c)
        return entry.value if entry is not None else None
    def text(r,c):
        return normalized_name(cell(r,c))
    warnings = []
    sections = {}
    for row in range(1,len(rows)+1):
        heading = text(row,1)
        for code,name,prefix in GROUPS:
            if heading.startswith(prefix) and _fy(cell(row,1)) is not None:
                if code in sections:
                    raise LFOValidationError(f'Multiple {name} store sections were found; import is ambiguous.')
                sections[code] = (row, _fy(cell(row,1)))
    if set(sections) != {g[0] for g in GROUPS}:
        raise LFOValidationError('The LFO sheet must contain Reliance, Lulu and Nilkamal financial-year store sections.')
    fys = {x[1] for x in sections.values()}
    if len(fys) != 1:
        raise LFOValidationError('The distributor store sections use different financial years.')
    fy = fys.pop()
    monetary = _MonetaryCells(workbook,worksheet)
    payload = {'fy':fy, 'source_id':SOURCE_ID,'sheet_name':SHEET,
               'distributors':[{'code':c,'name':n} for c,n,p in GROUPS],
               'stores':[], 'sales':[], 'weights':[], 'branchPlans':[], 'warnings':warnings}
    detail_totals = {}
    for code,name,_ in GROUPS:
        title_row = sections[code][0]
        header = title_row+1
        if text(header,2) != 'soh':
            raise LFOValidationError(f'{name} store monthly header was not found below its section heading.')
        dates = [_month_date(cell(header,col),fy,m,warnings)
                 for col,m in zip(range(3,15),MONTH_NUMBERS)]
        if text(header,17) not in ('grand total','total'):
            raise LFOValidationError(f'{name} store total column was not recognized.')
        seen = set()
        found_total = False
        for row in range(header+1,min(header+501,len(rows)+1)):
            key = text(row,1)
            if key in ('grand total','total'):
                found_total = True
                break
            if not key:
                if any(cell(row,c) is not None for c in range(3,15)):
                    raise LFOValidationError(f'{name} contains monthly sales without a store name.')
                continue
            if key in seen:
                raise LFOValidationError(f'{name} has duplicate normalized store names; correct the sheet before import.')
            seen.add(key)
            is_aggregate = key == 'others'
            store = {'distributorCode':code,'name':' '.join(str(cell(row,1)).split()),
                     'key':key,'isAggregate':is_aggregate,'sourceRow':row}
            payload['stores'].append(store)
            for col,month in zip(range(3,15),dates):
                amount = monetary(source_cell(row,col))
                payload['sales'].append({'distributorCode':code,'storeKey':key,
                    'period_month':month,'amount':amount,'sourceCell':f'{chr(64+col)}{row}'})
                detail_totals[(code,month)] = detail_totals.get((code,month),Decimal(0)) + (amount or Decimal(0))
        if not found_total or not seen:
            raise LFOValidationError(f'{name} store list must end with Total or Grand Total.')
        if not any(s['distributorCode']==code and not s['isAggregate'] for s in payload['stores']):
            raise LFOValidationError(f'{name} must have at least one named branch.')
    # Discover summary by exact labels and distributor headings, not row number.
    summary = []
    for row in range(2,len(rows)+1):
        if text(row,2)=='month' and all(text(row,c)=='monthly target' for c in (5,9,13)):
            if (text(row-1,3)=='lulu' and text(row-1,7)=='nilkamal homes'
                    and text(row-1,11)=='reliance'):
                summary.append(row)
    if len(summary)!=1:
        raise LFOValidationError('The LFO distributor monthly target summary is missing or ambiguous.')
    header = summary[0]
    source_difference = False
    for offset,expected in enumerate(MONTH_NUMBERS,1):
        row = header+offset
        month = _month_date(cell(row,2),fy,expected,warnings)
        for code,col in (('LULU',5),('NILKAMAL',9),('RELIANCE',13)):
            weight = monetary(source_cell(row,col))
            payload['weights'].append({'distributorCode':code,'period_month':month,
                                      'weight':weight,'sourceCell':f'{chr(64+col)}{row}'})
            summary_amount = monetary(source_cell(row,col+1))
            if summary_amount is not None and abs(summary_amount-detail_totals.get((code,month),Decimal(0)))>Decimal('0.01'):
                source_difference = True
    # Each FY branch plan is an independent table. Locate it by distributor,
    # year and the Store Name / Month / Year header, not fixed row positions.
    for code,name,prefix in GROUPS:
        candidates = []
        for row in range(1,len(rows)):
            for col in (3,7,11):
                title = text(row,col)
                if (title.startswith(prefix) and 'target' in title and _fy(cell(row,col))==fy
                        and [text(row+1,col+i) for i in range(3)]==['store name','month','year']):
                    candidates.append((row+1,col))
        if len(candidates)!=1:
            raise LFOValidationError(f'{name} needs one complete FY branch target table with Store Name, Month and Year columns.')
        header,col = candidates[0]
        area_end = min(start for start,_ in sections.values())
        if header>=area_end:
            raise LFOValidationError(f'{name} branch target table is outside the FY target area.')
        seen = set()
        terminated = False
        for row in range(header+1,min(header+501,area_end)):
            source_name = ' '.join(str(cell(row,col) or '').split())
            key = normalized_name(source_name)
            if key in ('grand total','total'):
                terminated = True
                break
            if not key:
                if any(cell(row,col+i) is not None for i in (1,2)):
                    raise LFOValidationError(f'{name} has a branch target amount without a branch name.')
                # A blank may terminate a table with no Grand Total, as in
                # Reliance. It must not hide later branches in this FY area.
                for later in range(row+1,area_end):
                    later_key = text(later,col)
                    if later_key in ('grand total','total'):
                        break
                    # The colour-band legend immediately above the sales
                    # section is not part of any branch target table.
                    band = [text(later,c) for c in (5,6,7)]
                    if (all(band) and all(re.fullmatch(r'[\d.,\sLl<>\-]+',v) for v in band)
                            and not any(cell(later,c) is not None for c in (1,2,3,4))):
                        continue
                    if any(cell(later,col+i) is not None for i in range(3)):
                        raise LFOValidationError(f'{name} branch target table has a blank gap before further entries; existing plans were preserved.')
                terminated = True
                break
            if key=='others' or key in seen:
                raise LFOValidationError(f'{name} branch targets contain an aggregate or duplicate branch name.')
            seen.add(key)
            monthly,yearly = monetary(source_cell(row,col+1)),monetary(source_cell(row,col+2))
            if monthly is not None and yearly is not None and monthly*12!=yearly:
                warnings.append(f'{name} / {source_name}: the Year target differs from twelve Month targets. Both source amounts are kept as entered.')
            payload['branchPlans'].append({'distributorCode':code,'fy':fy,'sourceName':source_name,
                'key':key,'monthlyTarget':monthly,'yearlyTarget':yearly,
                'sourceMonthCell':f'{chr(65+col)}{row}','sourceYearCell':f'{chr(66+col)}{row}'})
        if not terminated or not seen:
            raise LFOValidationError(f'{name} branch target table is empty or incomplete; existing plans were preserved.')
    if source_difference:
        warnings.append(TOTALS_NOTE)
    if monetary.recovered:
        addresses = list(dict.fromkeys(monetary.recovered))
        warnings.append(f"Recovered {len(addresses)} numeric amount cells shown with date formatting: {', '.join(addresses)}. "
                        'Original numeric values were used; change these cells to a number or currency format in the sheet.')
    warnings.append(SALES_BASIS_NOTE)
    digest = json.dumps(payload,sort_keys=True,default=str,ensure_ascii=False).encode('utf-8')
    payload['content_hash'] = hashlib.sha256(digest).hexdigest()
    return payload


def _llfo_month(value):
    """A month header in L-LFO Distributed: a date, 'Apr-2026', or a bare name
    such as 'April', 'June' or 'Sept'. Returns (month number, year or None)."""
    month, year = _source_month(value)
    if month is not None:
        return month, year
    word = str(value or '').strip().rstrip('.').casefold()
    number = LLFO_MONTH_WORDS.get(word[:3]) if len(word) >= 3 and word.isalpha() else None
    # A word is a month when its full name starts with it: 'Sept' and 'Aug'
    # qualify, 'Achieved' and 'Marketing' do not.
    if number and date(2000, number, 1).strftime('%B').casefold().startswith(word):
        return number, None
    return None, None


def parse_llfo_workbook(workbook, fy=None):
    """Validate L-LFO Distributed and return the same payload as parse_workbook.

    The sheet is one table: Group, Location, Primary and Secondary targets,
    then one secondary-sales column per month. Columns and rows are found from
    their header text, not fixed positions. Each Group becomes a distributor
    and each Location a store within it.

    Targets: each store's Secondary Target is its own monthly branch target;
    its year is that amount times twelve, because the sheet carries no Year
    cell. A group's distributor target is the sum of its stores' targets --
    the sheet has no separate group-level figure -- and is flagged when some
    of its stores have none.

    The sheet names months without a year, so the financial year comes from a
    heading such as '2026-27' if the sheet ever carries one, else from the
    caller. Guessing it from today's date would file a whole year of sales
    under the wrong year the moment the calendar turned over.
    """
    if LLFO_SHEET not in workbook.sheetnames:
        raise LFOValidationError('The workbook has no L-LFO Distributed worksheet.')
    worksheet = workbook[LLFO_SHEET]
    rows = list(worksheet.iter_rows())
    def source_cell(r,c):
        return rows[r-1][c-1] if 0 < r <= len(rows) and 0 < c <= len(rows[r-1]) else None
    def cell(r,c):
        entry = source_cell(r,c)
        return entry.value if entry is not None else None
    def text(r,c):
        return normalized_name(cell(r,c))
    def width(r):
        return len(rows[r-1]) if 0 < r <= len(rows) else 0
    def address(r,c):
        entry = source_cell(r,c)
        return getattr(entry,'coordinate',None) or f'{chr(64+c)}{r}'
    warnings = []

    # ---- the header: 'Group' beside 'Location', within the first 30 rows
    header = group_col = None
    for r in range(1,min(30,len(rows))+1):
        for c in range(1,width(r)):
            if text(r,c)=='group' and text(r,c+1)=='location':
                if header is not None:
                    raise LFOValidationError('L-LFO Distributed has more than one Group / Location header; import is ambiguous.')
                header,group_col = r,c
    if header is None:
        raise LFOValidationError('L-LFO Distributed needs a Group / Location header row.')
    loc_col, sub = group_col+1, header+1

    # ---- the Secondary Target column: 'Secondary' above 'Target'
    secondary = [c for c in range(1,width(header)+1)
                 if text(header,c)=='secondary' and text(sub,c)=='target']
    if len(secondary)!=1:
        raise LFOValidationError('L-LFO Distributed needs exactly one Secondary / Target column.')
    target_col = secondary[0]

    # ---- the monthly block: starts under 'Secondary Achieved ...', runs Apr -> Mar
    starts = [c for c in range(1,width(header)+1) if text(header,c).startswith('secondary achieved')]
    if len(starts)!=1:
        raise LFOValidationError('L-LFO Distributed needs one Secondary Achieved monthly block.')
    month_cols, header_years = [], {}
    for index,c in enumerate(range(starts[0],starts[0]+12)):
        month,year = _llfo_month(cell(sub,c))
        if month is None:
            break
        if month != MONTH_NUMBERS[index]:
            raise LFOValidationError('L-LFO monthly columns must run from April in order, without gaps or duplicates.')
        month_cols.append((c,month))
        if year is not None:
            header_years[c] = (month,year)
    if not month_cols:
        raise LFOValidationError('L-LFO Distributed has no monthly sales columns beginning with April.')
    # A month column after the run would be sales this import silently skips.
    after = month_cols[-1][0]+1
    if len(month_cols)<12 and _llfo_month(cell(sub,after+1))[0] is not None:
        raise LFOValidationError(f'L-LFO monthly headers have a gap at {address(sub,after)}; fill or remove it before import.')

    # ---- the financial year. Text headings only: a real date cell such as
    # 2026-04-01 would otherwise read as an inconsistent '2026-04'.
    found = {y for r in range(1,sub+1) for c in range(1,width(r)+1)
             if isinstance(cell(r,c),str) for y in [_fy(cell(r,c))] if y is not None}
    if len(found)>1:
        raise LFOValidationError('L-LFO Distributed names more than one financial year in its header.')
    if found:
        fy = found.pop()
    elif fy is None:
        raise LFOValidationError('L-LFO Distributed shows months without a year. Set the L-LFO financial year '
                                 '(llfo_secondary_fy) or add a heading such as 2026-27 to the sheet.')
    else:
        fy = int(fy)
        warnings.append(f'L-LFO Distributed has no financial-year heading; its months are read as FY {fy}–{str(fy+1)[-2:]} from configuration.')
    if not 2000 <= fy <= 2098:
        raise LFOValidationError('The L-LFO financial year is out of range.')
    months = {c:date(fy if m>=4 else fy+1,m,1) for c,m in month_cols}
    for c,(m,year) in header_years.items():
        if year != months[c].year:
            raise LFOValidationError(f'{address(sub,c)} is dated {year}, outside financial year {fy}–{str(fy+1)[-2:]}.')

    monetary = _MonetaryCells(workbook,worksheet)
    def amount_at(r,c):
        try:
            return monetary(source_cell(r,c))
        except LFOValidationError as exc:
            raise LFOValidationError(f'{address(r,c)}: {exc.public_message}') from None

    payload = {'fy':fy,'source_id':SOURCE_ID,'sheet_name':LLFO_SHEET,'distributors':[],
               'stores':[],'sales':[],'weights':[],'branchPlans':[],'warnings':warnings}
    groups, seen, last_row = {}, {}, None
    col_totals = {c:Decimal(0) for c,_ in month_cols}
    def named(r):
        return (' '.join(str(cell(r,group_col) or '').split()),
                ' '.join(str(cell(r,loc_col) or '').split()))
    for r in range(sub+1,len(rows)+1):
        group, location = named(r)
        if not group and not location:
            last_row = r
            # A blank row inside the table would end the import there and
            # silently drop every store below it.
            for later in range(r+1,min(r+60,len(rows)+1)):
                if all(named(later)):
                    raise LFOValidationError(f'L-LFO Distributed has a blank row at {r} inside its store table; '
                                             f'stores from row {later} would be missed. Remove the blank row.')
            break
        if not group or not location:
            raise LFOValidationError(f'Row {r} of L-LFO Distributed needs both a Group and a Location.')
        code = group.upper()
        if code not in groups:
            groups[code] = {'code':code,'name':group,'stores':[],'targets':[]}
            payload['distributors'].append({'code':code,'name':group})
        key = normalized_name(location)
        if key in seen.setdefault(code,set()):
            raise LFOValidationError(f'{group} lists {location} twice; correct the sheet before import.')
        if key=='others':
            raise LFOValidationError(f'{group} has an Others row; L-LFO stores must be named branches.')
        seen[code].add(key)
        payload['stores'].append({'distributorCode':code,'name':location,'key':key,
                                  'isAggregate':False,'sourceRow':r})
        for c,_ in month_cols:
            amount = amount_at(r,c)
            col_totals[c] += amount or Decimal(0)
            payload['sales'].append({'distributorCode':code,'storeKey':key,'period_month':months[c],
                                     'amount':amount,'sourceCell':address(r,c)})
        monthly = amount_at(r,target_col)
        groups[code]['stores'].append(location)
        groups[code]['targets'].append((address(r,target_col),monthly))
        payload['branchPlans'].append({'distributorCode':code,'fy':fy,'sourceName':location,'key':key,
            'monthlyTarget':monthly,'yearlyTarget':None if monthly is None else monthly*12,
            'sourceMonthCell':address(r,target_col),'sourceYearCell':f'{address(r,target_col)}×12'})
    else:
        # The table ran to the sheet's last row: a legitimate end, just with
        # no totals row beneath to cross-check against.
        last_row = len(rows)+1
    if not payload['stores']:
        raise LFOValidationError('L-LFO Distributed lists no stores.')

    # ---- a group's target is the sum of its stores' own targets
    for code,g in groups.items():
        known = [(a,v) for a,v in g['targets'] if v is not None]
        missing = len(g['targets'])-len(known)
        weight = sum((v for _,v in known),Decimal(0)) if known else None
        cells = '+'.join(a for a,_ in known) if known else f"{g['targets'][0][0]} (blank)"
        if known and missing:
            warnings.append(f"{g['name']}: {missing} of {len(g['targets'])} branches have no Secondary Target, "
                            'so its distributor target covers only the branches that have one.')
        for m in MONTH_NUMBERS:
            payload['weights'].append({'distributorCode':code,'period_month':date(fy if m>=4 else fy+1,m,1),
                                       'weight':weight,'sourceCell':cells})

    # ---- the sheet's own totals row, where present, is a cross-check only
    for c,_ in month_cols:
        stated = cell(last_row,c)
        if isinstance(stated,(int,float,Decimal)) and not isinstance(stated,bool):
            if abs(Decimal(str(stated))-col_totals[c]) > Decimal('0.01'):
                warnings.append(TOTALS_NOTE)
                break
    if len(month_cols)<12:
        first,last = months[month_cols[0][0]],months[month_cols[-1][0]]
        warnings.append(f'L-LFO Distributed has monthly columns for {first:%B} to {last:%B} only; later months are unreported until their columns are added.')
    if monetary.recovered:
        addresses = list(dict.fromkeys(monetary.recovered))
        warnings.append(f"Recovered {len(addresses)} numeric amount cells shown with date formatting: {', '.join(addresses)}. "
                        'Original numeric values were used; change these cells to a number or currency format in the sheet.')
    warnings.append(LLFO_NOT_IMPORTED_NOTE)
    digest = json.dumps(payload,sort_keys=True,default=str,ensure_ascii=False).encode('utf-8')
    payload['content_hash'] = hashlib.sha256(digest).hexdigest()
    return payload


def match_branch_plan(distributor_code,source_name,store_lookup):
    """Lookup keys are normalized names within one distributor; no guessing."""
    key = normalized_name(source_name)
    if key in store_lookup:
        return store_lookup[key]
    alias = BRANCH_TARGET_ALIASES.get(distributor_code,{}).get(key)
    return store_lookup.get(alias) if alias else None


def _actor(cur, actor):
    name = str(actor or '').strip()
    if not name or len(name)>200:
        raise LFOValidationError('An authenticated actor is required.')
    cur.execute("select set_config('app.lfo_actor',%s,true)",(name,))
    return name


def _channel_spec(channel):
    if channel not in CHANNELS:
        raise LFOValidationError('Unknown secondary sales channel.')
    return CHANNELS[channel]


def import_workbook(workbook, conn, actor='sheet-sync', channel='LFO', fy=None):
    """Atomic, idempotent upsert. Manual amount wins over future source updates.

    No source absence deletes stores or monthly sales. A blank replacing a prior
    value is reported and preserved so a failed formula cannot erase sales.
    channel picks the sheet and its parser; fy is the fallback financial year
    for L-LFO Distributed when the sheet names none.
    """
    _channel_spec(channel)
    if not re.fullmatch(r'[A-Za-z0-9_-]{20,200}', SOURCE_ID):
        raise LFOValidationError('Set SALES_SOURCE_ID to your own workbook ID before importing secondary sales.')
    data = parse_workbook(workbook) if channel=='LFO' else parse_llfo_workbook(workbook, fy=fy)
    source_id, sheet = data.get('source_id',SOURCE_ID), data.get('sheet_name',SHEET)
    counts = {'inserted':0,'updated':0,'unchanged':0,'manual_preserved':0,
              'stores_inserted':0,'weights_updated':0,'branch_targets_changed':0}
    warnings = list(data['warnings'])
    with conn:
        with conn.cursor() as cur:
            _actor(cur,actor)
            cur.execute("select pg_advisory_xact_lock(hashtext('lfo-secondary-import'))")
            # The second column reports whether migration 14 is applied. Without
            # it LFO keeps its original single-channel statements and L-LFO
            # refuses, rolling this whole transaction back.
            cur.execute("""insert into lfo.import_run(source_id,sheet_name,content_hash,actor) values(%s,%s,%s,%s)
                returning import_id,(select exists(select 1 from information_schema.columns
                where table_schema='lfo' and table_name='distributor' and column_name='channel'))""",
                        (source_id,sheet,data['content_hash'],actor))
            import_id,*migrated = cur.fetchone()
            has_channel = bool(migrated and migrated[0])
            if channel!='LFO' and not has_channel:
                raise LFOValidationError('Apply database migration 14_llfo_secondary.sql before importing L-LFO Distributed.')
            distributors = {}
            for d in data['distributors']:
                if has_channel:
                    cur.execute('insert into lfo.distributor(channel,code,name) values(%s,%s,%s) on conflict(channel,code) do nothing',
                                (channel,d['code'],d['name']))
                    cur.execute('select distributor_id from lfo.distributor where channel=%s and code=%s',(channel,d['code']))
                else:
                    cur.execute('insert into lfo.distributor(code,name) values(%s,%s) on conflict(code) do nothing',(d['code'],d['name']))
                    cur.execute('select distributor_id from lfo.distributor where code=%s',(d['code'],))
                distributors[d['code']] = cur.fetchone()[0]
            stores = {}
            for store in data['stores']:
                did = distributors[store['distributorCode']]
                cur.execute('insert into lfo.store(distributor_id,name,normalized_name,is_aggregate) values(%s,%s,%s,%s) on conflict(distributor_id,normalized_name) do nothing returning store_id',
                            (did,store['name'],store['key'],store['isAggregate']))
                created = cur.fetchone()
                if created:
                    counts['stores_inserted'] += 1
                cur.execute('select store_id,is_aggregate from lfo.store where distributor_id=%s and normalized_name=%s',(did,store['key']))
                sid,aggregate = cur.fetchone()
                if aggregate != store['isAggregate']:
                    raise LFOValidationError('A source store conflicts with an existing aggregate bucket.')
                stores[(store['distributorCode'],store['key'])] = sid
            # A group and year saved in the Target planner belongs to the planner:
            # this sheet's targets for it are skipped, never merged or overwritten.
            fy = data['fy']
            cur.execute('''select distributor_id from lfo.branch_target_plan where fy=%s and source_id=%s and distributor_id=any(%s)
                union select distributor_id from lfo.distributor_weight where source_id=%s
                and period_month>=%s and period_month<%s and distributor_id=any(%s)''',
                (fy,PLANNER_SOURCE,list(distributors.values()),PLANNER_SOURCE,date(fy,4,1),date(fy+1,4,1),
                 list(distributors.values())))
            planned = {row[0] for row in cur.fetchall()}
            for d in data['distributors']:
                if distributors[d['code']] in planned:
                    warnings.append(PLANNED_NOTE.format(name=d['name'],fy=fy,yy=str(fy+1)[2:]))
            if 'branchPlans' in data:
                # Parse validated every group's complete source table before
                # starting this transaction. Preserve all historical FY plans.
                for code,did in distributors.items():
                    if did in planned:
                        continue
                    plans = [p for p in data['branchPlans'] if p['distributorCode']==code]
                    cur.execute('select normalized_name,store_id from lfo.store where distributor_id=%s and not is_aggregate',(did,))
                    lookup = dict(cur.fetchall())
                    matched = set()
                    for plan in plans:
                        plan['storeId'] = match_branch_plan(code,plan['sourceName'],lookup)
                        sid = plan['storeId']
                        if sid is not None:
                            if sid in matched:
                                raise LFOValidationError('Two source branch plans match the same store. Correct the source names before importing.')
                            matched.add(sid)
                        else:
                            warnings.append(f"{plan['sourceName']} has a branch target but no confirmed matching store under {code}.")
                    cur.execute('''update lfo.branch_target_plan set is_active=false,updated_at=now(),updated_by=%s,import_id=%s
                        where distributor_id=%s and fy=%s and source_id=%s and is_active
                        and not(normalized_name=any(%s))''',
                        (actor,import_id,did,data['fy'],source_id,[p['key'] for p in plans]))
                    counts['branch_targets_changed'] += cur.rowcount
                    for plan in plans:
                        cur.execute('''insert into lfo.branch_target_plan(distributor_id,fy,source_name,normalized_name,
                            monthly_target,yearly_target,store_id,is_active,source_id,sheet_name,
                            source_month_cell,source_year_cell,import_id,updated_by)
                            values(%s,%s,%s,%s,%s,%s,%s,true,%s,%s,%s,%s,%s,%s)
                            on conflict(distributor_id,fy,normalized_name) do update set
                            source_name=excluded.source_name,monthly_target=excluded.monthly_target,yearly_target=excluded.yearly_target,
                            store_id=excluded.store_id,is_active=true,source_id=excluded.source_id,sheet_name=excluded.sheet_name,
                            source_month_cell=excluded.source_month_cell,source_year_cell=excluded.source_year_cell,
                            import_id=excluded.import_id,updated_by=excluded.updated_by,updated_at=now()
                            where lfo.branch_target_plan.source_id<>'dashboard'
                            and (lfo.branch_target_plan.source_name,lfo.branch_target_plan.monthly_target,
                            lfo.branch_target_plan.yearly_target,lfo.branch_target_plan.store_id,lfo.branch_target_plan.is_active,
                            lfo.branch_target_plan.source_month_cell,lfo.branch_target_plan.source_year_cell)
                            is distinct from (excluded.source_name,excluded.monthly_target,excluded.yearly_target,excluded.store_id,
                            true,excluded.source_month_cell,excluded.source_year_cell)''',
                            (did,plan['fy'],plan['sourceName'],plan['key'],plan['monthlyTarget'],plan['yearlyTarget'],
                             plan['storeId'],source_id,sheet,plan['sourceMonthCell'],plan['sourceYearCell'],import_id,actor))
                        counts['branch_targets_changed'] += cur.rowcount
            for item in data['weights']:
                if distributors[item['distributorCode']] in planned:
                    continue
                if item['weight'] is None:
                    warnings.append(f"No monthly target is entered in {item['sourceCell']}; that distributor's target is unavailable for the month.")
                cur.execute('''insert into lfo.distributor_weight(distributor_id,period_month,weight,source_id,source_cell,import_id,updated_by)
                    values(%s,%s,%s,%s,%s,%s,%s)
                    on conflict(distributor_id,period_month) do update set weight=excluded.weight,
                    source_id=excluded.source_id,source_cell=excluded.source_cell,import_id=excluded.import_id,
                    updated_by=excluded.updated_by,updated_at=now()
                    where lfo.distributor_weight.source_id<>'dashboard'
                    and (lfo.distributor_weight.weight,lfo.distributor_weight.source_cell)
                    is distinct from (excluded.weight,excluded.source_cell)''',
                    (distributors[item['distributorCode']],item['period_month'],item['weight'],source_id,item['sourceCell'],import_id,actor))
                counts['weights_updated'] += cur.rowcount
            for item in data['sales']:
                sid = stores[(item['distributorCode'],item['storeKey'])]
                mon,amount = item['period_month'],item['amount']
                cur.execute('select amount,origin,source_amount,source_cell from lfo.monthly_sale where store_id=%s and period_month=%s for update',(sid,mon))
                old = cur.fetchone()
                if amount is None:
                    if old and old[2] is not None:
                        warnings.append(f"Blank source cell {item['sourceCell']} did not erase its existing monthly sale.")
                    continue
                if old is None:
                    # A concurrent form create may win while this import is in
                    # progress. ON CONFLICT preserves it; next sync adds provenance.
                    cur.execute('''insert into lfo.monthly_sale(store_id,period_month,amount,origin,source_amount,source_id,source_cell,import_id,updated_by)
                        values(%s,%s,%s,'sheet',%s,%s,%s,%s,%s) on conflict(store_id,period_month) do nothing''',
                        (sid,mon,amount,amount,source_id,item['sourceCell'],import_id,actor))
                    counts['inserted'] += cur.rowcount
                    continue
                target_amount = old[0] if old[1]=='manual' else amount
                if old[1]=='manual':
                    counts['manual_preserved'] += 1
                if (old[0],old[2],old[3]) == (target_amount,amount,item['sourceCell']):
                    counts['unchanged'] += 1
                    continue
                cur.execute('''update lfo.monthly_sale set amount=%s,source_amount=%s,source_id=%s,source_cell=%s,
                    import_id=%s,updated_by=%s,updated_at=now(),revision=revision+1 where store_id=%s and period_month=%s''',
                    (target_amount,amount,source_id,item['sourceCell'],import_id,actor,sid,mon))
                counts['updated'] += 1
            warnings = list(dict.fromkeys(warnings))
            cur.execute('update lfo.import_run set warnings=%s,counters=%s where import_id=%s',
                        (Json(warnings),Json(counts),import_id))
    return {**counts,'changed':any(counts[k] for k in ('inserted','updated','stores_inserted','weights_updated','branch_targets_changed')),
            'warnings':warnings,'fy':data['fy'],'import_id':import_id}


def allocate_cents(total, weights):
    """Largest-remainder allocation, stable keys break equal-remainder ties."""
    amount = money(total)
    values = {key:money(weight) for key,weight in weights.items()}
    if not values:
        if amount:
            raise LFOValidationError('Cannot allocate a non-zero target without recipients.')
        return {}
    denominator = sum(values.values(),Decimal(0))
    if not denominator:
        values = {key:Decimal(1) for key in values}
        denominator = Decimal(len(values))
    cents = int(amount*100)
    exact = {k:Decimal(cents)*w/denominator for k,w in values.items()}
    allocated = {k:int(v.to_integral_value(rounding=ROUND_FLOOR)) for k,v in exact.items()}
    remaining = cents-sum(allocated.values())
    order = sorted(values,key=lambda k:(-(exact[k]-allocated[k]),str(k)))
    for key in order[:remaining]:
        allocated[key] += 1
    return {key:Decimal(value)/100 for key,value in allocated.items()}


def _friendly_warnings(warnings):
    """Keep existing import history readable after the target basis changed."""
    result = []
    for warning in warnings:
        if warning.startswith('Secondary amounts retain the sheet sales basis'):
            warning = SALES_BASIS_NOTE
        elif warning.startswith('Distributor summary sales differ from the branch monthly cells'):
            warning = TOTALS_NOTE
        elif 'source dates interpreted as' in warning:
            match = re.search(r'interpreted as (\d{4}) from FY (\d{4})-(\d{2})',warning)
            if match:
                year,start,end = match.groups()
                warning = f'January to March dates use {year} to match financial year {start}–{end}.'
        elif warning.startswith('No distributor weight in '):
            source_cell = warning.split(' in ',1)[1].split(';',1)[0]
            warning = f'No monthly target was entered in {source_cell}. Sync the sheet to refresh missing targets.'
        result.append(warning)
    return list(dict.fromkeys(result))


def fetch_dashboard(conn, channel_targets=None, channel='LFO'):
    """JSON payload consumed by dashboard and monthly-entry pages.

    channel is 'LFO' or 'L-LFO'. Each payload holds only that channel's
    distributors and everything beneath them, so neither page nor either
    monthly-entry form can see or change the other channel's stores.

    {available, distributors:[{id,name}], stores:[{id,distributorId,name,isAggregate}],
     sales:[{storeId,month,amount,revision,origin}],
     targets:[{storeId,distributorId,month,amount}],
     distributorTargets:[{distributorId,month,amount}],
     channelTargets:[{month,amount}], warnings:[str]}
    Month labels are Mon-YYYY. All monetary amounts are JSON numbers in INR.
    Distributor monthly targets are the sheet's separate secondary amounts.
    Store targets retain their branch table amounts without rescaling. channelTargets is a
    compatibility name for their monthly sum, also exposed as secondaryTargets.
    channel_targets is accepted for older callers but deliberately ignored.
    """
    spec = _channel_spec(channel)
    with conn.cursor() as cur:
        cur.execute("""select to_regclass('lfo.monthly_sale'),to_regclass('lfo.branch_target_plan'),
            exists(select 1 from information_schema.columns
                   where table_schema='lfo' and table_name='distributor' and column_name='channel')""")
        installed,plans_installed,*migrated = cur.fetchone()
        has_channel = bool(migrated and migrated[0])
        if not installed or (channel!='LFO' and not has_channel):
            return {'available':False,'distributors':[],'stores':[],'sales':[],
                    'targets':[],'distributorTargets':[],'channelTargets':[], 'secondaryTargets':[],
                    'targetBasis':spec['targetBasis'],'storeTargetBasis':spec['storeTargetBasis'],'branchTargetPlans':[],
                    'warnings':[f'{channel} store data has not been installed.']}
        # Without migration 14 every row is LFO, so the original unscoped
        # statements are exactly right; with it, each read is scoped to the channel.
        scope, args = ('where d.channel=%s',(channel,)) if has_channel else ('',())
        cur.execute(f'select d.distributor_id,d.name,d.code from lfo.distributor d {scope} order by d.code',args)
        distributors = [{'id':r[0],'name':r[1],'code':r[2]} for r in cur.fetchall()]
        cur.execute(f'''select s.store_id,s.distributor_id,s.name,s.is_aggregate from lfo.store s
            join lfo.distributor d using(distributor_id) {scope} order by s.distributor_id,s.normalized_name''',args)
        stores = [{'id':r[0],'distributorId':r[1],'name':r[2],'isAggregate':r[3]} for r in cur.fetchall()]
        cur.execute(f'''select m.store_id,m.period_month,m.amount,m.revision,m.origin,m.source_amount from lfo.monthly_sale m
            join lfo.store s using(store_id) join lfo.distributor d using(distributor_id) {scope}
            order by m.period_month,m.store_id''',args)
        sales = [{'storeId':r[0],'month':r[1].strftime('%b-%Y'),'amount':float(r[2]),'revision':r[3],
                  'origin':r[4],'sourceAmount':float(r[5]) if r[5] is not None else None} for r in cur.fetchall()]
        cur.execute(f'''select w.distributor_id,w.period_month,w.weight from lfo.distributor_weight w
            join lfo.distributor d using(distributor_id) {scope}''',args)
        weights = {(r[0],r[1]):r[2] for r in cur.fetchall()}
        plans = []
        if plans_installed:
            cur.execute(f'''select p.distributor_id,p.store_id,p.source_name,p.fy,p.monthly_target,p.yearly_target,
                p.source_month_cell,p.source_year_cell from lfo.branch_target_plan p
                join lfo.distributor d using(distributor_id) where p.is_active {scope.replace('where','and')}
                order by p.fy,p.distributor_id,p.normalized_name''',args)
            plans = [{'distributorId':r[0],'storeId':r[1],'sourceName':r[2],'fy':r[3],
                      'monthlyTarget':float(r[4]) if r[4] is not None else None,
                      'yearlyTarget':float(r[5]) if r[5] is not None else None,
                      'sourceMonthCell':r[6],'sourceYearCell':r[7]} for r in cur.fetchall()]
        # Each channel's warnings come from its own sheet's latest import. A
        # channel set up in the Target planner has no sheet and no import.
        imported = None
        if spec['sheet'] is None:
            warnings = []
        else:
            if has_channel:
                cur.execute('select warnings,imported_at from lfo.import_run where sheet_name=%s order by import_id desc limit 1',
                            (spec['sheet'],))
            else:
                cur.execute('select warnings,imported_at from lfo.import_run order by import_id desc limit 1')
            imported = cur.fetchone()
            warnings = _friendly_warnings(imported[0]) if imported else [f"{spec['sheet']} has not been imported."]
    source = ({'sheet':spec['sheet'],'spreadsheetId':SOURCE_ID or None,
               'url':(f"https://docs.google.com/spreadsheets/d/{SOURCE_ID}/edit#gid={spec['gid']}"
                      if SOURCE_ID else None)}
              if spec['sheet'] else {'sheet':PLANNER_SHEET,'spreadsheetId':None,'url':'/targets/planner'})
    # A planner channel is available once it has a sub-channel, even before
    # its first branch is added.
    payload = {'available':bool(stores) or (spec['sheet'] is None and bool(distributors)),
               'distributors':distributors,'stores':stores,
               'sales':sales,'targets':[],'distributorTargets':[], 'channelTargets':[],
               'warnings':warnings,'lastImport':imported[1].isoformat() if imported else None,
               'targetBasis':spec['targetBasis'],
               'storeTargetBasis':spec['storeTargetBasis'],'branchTargetPlans':plans,
               'targetSource':{**source,'description':spec['description'],'storeSplit':spec['storeSplit']}}
    payload['secondaryTargets'] = payload['channelTargets']
    # A stored secondary target month defines its fiscal year, independently
    # of whether an active primary plan exists for that financial year.
    years = sorted({month.year if month.month>=4 else month.year-1 for _,month in weights}
                   | {plan['fy'] for plan in plans})
    payload['financialYears'] = years
    for fy in years:
        for monthnum in MONTH_NUMBERS:
            month = date(fy if monthnum>=4 else fy+1,monthnum,1)
            label = month.strftime('%b-%Y')
            known = {d['id']:money(weights[(d['id'],month)]) for d in distributors
                     if weights.get((d['id'],month)) is not None}
            missing = [d for d in distributors if d['id'] not in known]
            if missing:
                warnings.append(f"{label}: no monthly target is entered for {', '.join(d['name'] for d in missing)}.")
            if known:
                payload['channelTargets'].append({'month':label,'amount':float(sum(known.values(),Decimal(0))),
                    'complete':not missing,'missingDistributorIds':[d['id'] for d in missing]})
            for did,value in known.items():
                payload['distributorTargets'].append({'distributorId':did,'month':label,'amount':float(value)})
    if not plans_installed:
        warnings.append('Branch target plans have not been installed; no store targets are estimated.')
    for plan in plans:
        if plan['storeId'] is None or plan['monthlyTarget'] is None:
            continue
        for monthnum in MONTH_NUMBERS:
            month = date(plan['fy'] if monthnum>=4 else plan['fy']+1,monthnum,1)
            payload['targets'].append({'storeId':plan['storeId'],'distributorId':plan['distributorId'],
                                      'month':month.strftime('%b-%Y'),'amount':plan['monthlyTarget']})
    payload['warnings'] = list(dict.fromkeys(warnings))
    return payload


def save_monthly_sale(conn, store_id, period_month, amount, expected_revision, actor):
    """Create/update one monthly secondary amount with optimistic concurrency."""
    if type(store_id) is not int or store_id<=0:
        raise LFOValidationError('Select an existing LFO store.')
    if type(expected_revision) is not int or expected_revision<0:
        raise LFOValidationError('A valid record revision is required. Reload the page.')
    month,value = period(period_month),money(amount)
    with conn:
        with conn.cursor() as cur:
            _actor(cur,actor)
            cur.execute('select store_id from lfo.store where store_id=%s',(store_id,))
            if cur.fetchone() is None:
                raise LFOValidationError('The selected LFO store does not exist.')
            if expected_revision==0:
                cur.execute('''insert into lfo.monthly_sale(store_id,period_month,amount,origin,updated_by)
                    values(%s,%s,%s,'manual',%s) on conflict(store_id,period_month) do nothing returning revision''',
                    (store_id,month,value,actor))
            else:
                cur.execute('''update lfo.monthly_sale set amount=%s,origin='manual',revision=revision+1,
                    updated_at=now(),updated_by=%s where store_id=%s and period_month=%s and revision=%s returning revision''',
                    (value,actor,store_id,month,expected_revision))
            saved = cur.fetchone()
            if saved is None:
                raise LFOConflictError()
    return {'storeId':store_id,'periodMonth':month.isoformat(),'month':month.strftime('%b-%Y'),
            'amount':float(value),'revision':saved[0],'origin':'manual'}


def _distributor_in_channel(cur, distributor_id, channel):
    """Lock the distributor and confirm it belongs to channel (None: any).

    to_jsonb(d)->>'channel' is NULL rather than an error before migration 14,
    when every distributor is LFO. This is what stops a form for one channel
    writing under the other channel's distributor ID.
    """
    cur.execute("select distributor_id,to_jsonb(d)->>'channel' from lfo.distributor d "
                "where distributor_id=%s for key share",(distributor_id,))
    row = cur.fetchone()
    if row is None:
        return False
    found = (row[1] if len(row)>1 else None) or 'LFO'
    return channel is None or found==channel


def create_store(conn, distributor_id, name, actor, channel=None):
    """Add a named branch without altering primary sales or channel targets.

    Names match within a distributor after case and whitespace normalization.
    Store creation and its audit event commit together; a concurrent duplicate
    receives the same validation error as an already-existing name. Pass
    channel so a distributor from the other channel is refused.
    """
    if type(distributor_id) is not int or distributor_id<=0:
        raise LFOValidationError('Select an existing LFO distributor.')
    if not isinstance(name,str) or any(unicodedata.category(c) in ('Cc','Cf','Cs') for c in name):
        raise LFOValidationError('Enter a branch name without control characters.')
    display_name = ' '.join(name.split())
    if not display_name or len(display_name)>200:
        raise LFOValidationError('Enter a branch name between 1 and 200 characters.')
    key = normalized_name(display_name)
    if key=='others':
        raise LFOValidationError('Others is reserved for aggregate sales. Enter the branch or retailer name.')
    with conn:
        with conn.cursor() as cur:
            actor = _actor(cur,actor)
            if not _distributor_in_channel(cur,distributor_id,channel):
                raise LFOValidationError(f'The selected distributor / group was not found in the {channel} channel.'
                                         if channel else 'The selected LFO distributor does not exist.')
            cur.execute('''insert into lfo.store(distributor_id,name,normalized_name,is_aggregate)
                values(%s,%s,%s,false) on conflict(distributor_id,normalized_name) do nothing returning store_id''',
                (distributor_id,display_name,key))
            created = cur.fetchone()
            if created is None:
                raise LFOValidationError('A branch with this name already exists under the selected distributor.')
            store_id = created[0]
            cur.execute('''insert into lfo.audit_log(actor,table_name,operation,before_row,after_row)
                values(%s,'store','INSERT',null,%s)''',
                (actor,Json({'store_id':store_id,'distributor_id':distributor_id,'name':display_name,
                             'normalized_name':key,'is_aggregate':False})))
    return {'id':store_id,'distributorId':distributor_id,'name':display_name,'isAggregate':False}


def save_monthly_sales(conn, distributor_id, period_month, entries, actor, *,
                       allow_create=True, allow_edit=True, channel=None):
    """Atomically save changed monthly amounts for one distributor.

    entries is a list of {store_id, amount, expected_revision}; zero revision
    denotes a new record. Blank amounts are invalid and should be excluded by
    the caller. Permissions are checked against the locked database records,
    never inferred solely from the submitted revision. Every store ownership,
    row revision and permission is validated before any monthly write.
    """
    if type(distributor_id) is not int or distributor_id<=0:
        raise LFOValidationError('Select an existing LFO distributor.')
    if type(allow_create) is not bool or type(allow_edit) is not bool:
        raise LFOValidationError('Valid monthly sales permissions are required.')
    if not isinstance(entries,list) or not 1<=len(entries)<=500:
        raise LFOValidationError('Submit between 1 and 500 changed monthly sales amounts.')
    month = period(period_month)
    parsed = {}
    for item in entries:
        if not isinstance(item,dict):
            raise LFOValidationError('Each monthly sale must identify its store, amount and revision.')
        store_id,revision = item.get('store_id'),item.get('expected_revision')
        if type(store_id) is not int or store_id<=0:
            raise LFOValidationError('Select an existing LFO store for each monthly amount.')
        if store_id in parsed:
            raise LFOValidationError('A store can appear only once in a monthly sales batch.')
        if type(revision) is not int or revision<0:
            raise LFOValidationError('A valid record revision is required for every store. Reload the page.')
        parsed[store_id] = (money(item.get('amount')),revision)
    ordered_ids = sorted(parsed)
    saved = []
    with conn:
        with conn.cursor() as cur:
            actor = _actor(cur,actor)
            if not _distributor_in_channel(cur,distributor_id,channel):
                raise LFOValidationError('The selected LFO distributor does not exist.')
            # Stable store locks prevent an ownership change midway through a
            # batch, and impose one lock order for overlapping batch requests.
            cur.execute('''select store_id,distributor_id from lfo.store where store_id=any(%s)
                order by store_id for share''',(ordered_ids,))
            ownership = dict(cur.fetchall())
            if len(ownership)!=len(ordered_ids) or any(ownership[sid]!=distributor_id for sid in ordered_ids):
                raise LFOValidationError('Every selected store must belong to the selected LFO distributor.')
            cur.execute('''select store_id,revision from lfo.monthly_sale
                where store_id=any(%s) and period_month=%s order by store_id for update''',(ordered_ids,month))
            revisions = dict(cur.fetchall())
            for store_id in ordered_ids:
                _,expected_revision = parsed[store_id]
                current_revision = revisions.get(store_id)
                if current_revision is None:
                    if not allow_create:
                        raise LFOValidationError('You do not have access to create monthly store sales.')
                    if expected_revision!=0:
                        raise LFOConflictError()
                else:
                    if not allow_edit:
                        raise LFOValidationError('You do not have access to edit existing monthly store sales.')
                    if expected_revision!=current_revision:
                        raise LFOConflictError()
            for store_id in ordered_ids:
                value,expected_revision = parsed[store_id]
                if expected_revision==0:
                    # Another writer may have inserted this previously absent
                    # row. Never promote a create into an update on conflict.
                    cur.execute('''insert into lfo.monthly_sale(store_id,period_month,amount,origin,updated_by)
                        values(%s,%s,%s,'manual',%s) on conflict(store_id,period_month) do nothing returning revision''',
                        (store_id,month,value,actor))
                else:
                    cur.execute('''update lfo.monthly_sale set amount=%s,origin='manual',revision=revision+1,
                        updated_at=now(),updated_by=%s where store_id=%s and period_month=%s and revision=%s returning revision''',
                        (value,actor,store_id,month,expected_revision))
                result = cur.fetchone()
                if result is None:
                    raise LFOConflictError()
                saved.append({'storeId':store_id,'periodMonth':month.isoformat(),
                              'month':month.strftime('%b-%Y'),'amount':float(value),
                              'revision':result[0],'origin':'manual'})
    return {'saved':len(saved),'sales':saved}


# ------------------------------------------------------------ Target planner

def planner_installed(cur, channel):
    """True when the database accepts this channel's sub-channels (migration 15
    for D&R-S, D&R-N and Saloon; LFO and L-LFO need only migration 14)."""
    cur.execute("select pg_get_constraintdef(oid) from pg_constraint where conname='distributor_channel_check'")
    row = cur.fetchone()
    return bool(row) and f"'{channel}'" in row[0]


def create_distributor(conn, channel, name, actor):
    """Add a sub-channel / group to a channel from the Target planner.

    The code is the name in capitals, unique within the channel, as the sheet
    imports make it. Returns {'id','name','code'}; an existing name is refused.
    """
    _channel_spec(channel)
    if not isinstance(name,str) or any(unicodedata.category(c) in ('Cc','Cf','Cs') for c in name):
        raise LFOValidationError('Enter a sub-channel name without control characters.')
    display_name = ' '.join(name.split())
    if not display_name or len(display_name)>120:
        raise LFOValidationError('Enter a sub-channel name between 1 and 120 characters.')
    code = display_name.upper()
    with conn:
        with conn.cursor() as cur:
            actor = _actor(cur,actor)
            if not planner_installed(cur,channel):
                raise LFOValidationError(f'{channel} sub-channels need the database update. Run apply-update.cmd on the '
                                         'computer that hosts the database, then try again.')
            cur.execute('select 1 from lfo.distributor where channel=%s and upper(code)=%s',(channel,code))
            if cur.fetchone():
                raise LFOValidationError(f'{display_name} already exists in the {channel} channel.')
            cur.execute('''insert into lfo.distributor(channel,code,name) values(%s,%s,%s)
                on conflict(channel,code) do nothing returning distributor_id''',(channel,code,display_name))
            created = cur.fetchone()
            if created is None:
                raise LFOValidationError(f'{display_name} already exists in the {channel} channel.')
            cur.execute('''insert into lfo.audit_log(actor,table_name,operation,before_row,after_row)
                values(%s,'distributor','INSERT',null,%s)''',
                (actor,Json({'distributor_id':created[0],'channel':channel,'code':code,'name':display_name})))
    return {'id':created[0],'name':display_name,'code':code}


def fy_months(fy):
    """The twelve first-of-month dates of financial year fy, April to March."""
    return [date(fy if m>=4 else fy+1,m,1) for m in MONTH_NUMBERS]


def _planner_rows(cur, channel, fy):
    """Everything the planner shows or compares for one channel, for fy and fy-1."""
    cur.execute('select distributor_id,name,code from lfo.distributor where channel=%s order by name',(channel,))
    groups = [{'id':r[0],'name':r[1],'code':r[2]} for r in cur.fetchall()]
    ids = [g['id'] for g in groups]
    cur.execute('''select store_id,distributor_id,name,normalized_name,is_aggregate from lfo.store
        where distributor_id=any(%s) order by distributor_id,normalized_name''',(ids,))
    stores = [{'id':r[0],'distributorId':r[1],'name':r[2],'key':r[3],'isAggregate':r[4]} for r in cur.fetchall()]
    cur.execute('''select distributor_id,period_month,weight,source_id from lfo.distributor_weight
        where distributor_id=any(%s) and period_month>=%s and period_month<%s''',
        (ids,date(fy-1,4,1),date(fy+1,4,1)))
    weights = {(r[0],r[1]):(r[2],r[3]) for r in cur.fetchall()}
    cur.execute('''select plan_id,distributor_id,fy,store_id,source_name,normalized_name,monthly_target,yearly_target,source_id
        from lfo.branch_target_plan where distributor_id=any(%s) and fy in (%s,%s) and is_active
        order by plan_id''',(ids,fy-1,fy))
    plans = [{'planId':r[0],'distributorId':r[1],'fy':r[2],'storeId':r[3],'sourceName':r[4],'key':r[5],
              'monthly':r[6],'yearly':r[7],'source':r[8]} for r in cur.fetchall()]
    return groups, stores, weights, plans


def _state_token(groups, stores, weights, plans, fy):
    """A digest of the stored year, so a save made from an older page is refused."""
    months = set(fy_months(fy))
    rows = [sorted((g['id'],g['name']) for g in groups), sorted((s['id'],s['distributorId'],s['name']) for s in stores),
            sorted((k[0],k[1].isoformat(),str(v[0]),v[1]) for k,v in weights.items() if k[1] in months),
            sorted((p['planId'],p['storeId'],p['key'],str(p['monthly']),str(p['yearly']),p['source'])
                   for p in plans if p['fy']==fy)]
    return hashlib.sha256(json.dumps(rows,default=str).encode('utf-8')).hexdigest()[:32]


def _same(a, b):
    if a is None or b is None:
        return a is None and b is None
    return Decimal(str(a)) == Decimal(str(b))


def planner_secondary(conn, channel, fy):
    """One channel's secondary targets for fy, with fy-1 for reference.

    {installed, token, groups:[{id,name,source,months:[12],lastMonths:[12],
      branches:[{id,name,monthly,lastMonthly}], unmatched:[{name,monthly}]}]}
    Amounts are Decimal, or None when not set. source is 'planner', 'sheet'
    or None: where that group's targets for the year come from.
    """
    _channel_spec(channel)
    with conn.cursor() as cur:
        installed = planner_installed(cur,channel)
        groups, stores, weights, plans = _planner_rows(cur,channel,fy)
    this, last = fy_months(fy), fy_months(fy-1)
    result = []
    for g in groups:
        mine = [p for p in plans if p['distributorId']==g['id']]
        current = [p for p in mine if p['fy']==fy]
        sources = {weights[(g['id'],m)][1] for m in this if (g['id'],m) in weights}
        sources |= {p['source'] for p in current}
        source = 'planner' if PLANNER_SOURCE in sources else 'sheet' if sources else None
        by_store = {p['storeId']:p for p in current if p['storeId'] is not None}
        last_store = {p['storeId']:p for p in mine if p['fy']==fy-1 and p['storeId'] is not None}
        result.append({
            'id':g['id'],'name':g['name'],'source':source,
            'months':[weights.get((g['id'],m),(None,None))[0] for m in this],
            'lastMonths':[weights.get((g['id'],m),(None,None))[0] for m in last],
            'branches':[{'id':s['id'],'name':s['name'],
                         'monthly':by_store[s['id']]['monthly'] if s['id'] in by_store else None,
                         'lastMonthly':last_store[s['id']]['monthly'] if s['id'] in last_store else None}
                        for s in stores if s['distributorId']==g['id'] and not s['isAggregate']],
            'unmatched':[{'name':p['sourceName'],'monthly':p['monthly']} for p in current if p['storeId'] is None],
        })
    return {'installed':installed,'token':_state_token(groups,stores,weights,plans,fy),'groups':result}


def save_secondary_targets(conn, channel, fy, submitted, token, actor, allow_clear=False):
    """Save the planner's targets for one channel and year in one transaction.

    submitted is {distributor_id: {'months':[12 x Decimal|None],
    'branches':{store_id: Decimal|None}}}. A group whose values match what is
    stored is left alone. A group that changed is saved whole as a planner
    group for that year: its twelve monthly targets and one monthly target per
    branch (the branch's year is twelve of them). Imported plans for that group
    and year are set inactive, never deleted, and later sheet imports skip it.

    token must equal planner_secondary's for the stored year, or nothing is
    saved (LFOConflictError). Clearing a stored target needs allow_clear.
    Returns counters.
    """
    _channel_spec(channel)
    this = fy_months(fy)
    counts = {'groups':0,'months':0,'branches':0}
    with conn:
        with conn.cursor() as cur:
            actor = _actor(cur,actor)
            # One writer at a time for secondary targets: sheet imports and planner saves.
            cur.execute("select pg_advisory_xact_lock(hashtext('lfo-secondary-import'))")
            groups, stores, weights, plans = _planner_rows(cur,channel,fy)
            if token != _state_token(groups,stores,weights,plans,fy):
                raise LFOConflictError()
            known = {g['id'] for g in groups}
            run_id = None
            for did, values in submitted.items():
                if did not in known:
                    raise LFOValidationError(f'A sub-channel on this page is not in the {channel} channel. Reload the planner.')
                months = list(values.get('months') or [None]*12)
                if len(months)!=12:
                    raise LFOValidationError('Each sub-channel needs twelve monthly targets.')
                named = {s['id']:s for s in stores if s['distributorId']==did and not s['isAggregate']}
                active = [p for p in plans if p['distributorId']==did and p['fy']==fy]
                by_store = {p['storeId']:p for p in active if p['storeId'] is not None}
                stored_branches = {sid:(by_store[sid]['monthly'] if sid in by_store else None) for sid in named}
                branch_values = dict(stored_branches)
                for sid, amount in (values.get('branches') or {}).items():
                    if sid not in named:
                        raise LFOValidationError('A branch on this page does not belong to its sub-channel. Reload the planner.')
                    branch_values[sid] = amount
                stored_months = [weights.get((did,m),(None,None))[0] for m in this]
                if (all(_same(a,b) for a,b in zip(months,stored_months))
                        and all(_same(branch_values[s],stored_branches[s]) for s in named)):
                    continue
                cleared = any(new is None and old is not None for new,old in zip(months,stored_months)) or any(
                    branch_values[s] is None and stored_branches[s] is not None for s in named)
                if cleared and not allow_clear:
                    raise LFOValidationError('Clearing a saved target needs Targets Delete permission. Enter 0 instead, or ask an administrator.')
                counts['groups'] += 1
                for month, amount in zip(this, months):
                    if amount is None and (did,month) not in weights:
                        continue
                    cur.execute('''insert into lfo.distributor_weight(distributor_id,period_month,weight,source_id,source_cell,import_id,updated_by)
                        values(%s,%s,%s,%s,%s,null,%s) on conflict(distributor_id,period_month) do update set
                        weight=excluded.weight,source_id=excluded.source_id,source_cell=excluded.source_cell,
                        import_id=null,updated_by=excluded.updated_by,updated_at=now()''',
                        (did,month,amount,PLANNER_SOURCE,PLANNER_SHEET,actor))
                    counts['months'] += 1
                imported = [p['planId'] for p in active if p['source']!=PLANNER_SOURCE]
                if imported:
                    cur.execute('''update lfo.branch_target_plan set is_active=false,updated_at=now(),updated_by=%s
                        where plan_id=any(%s)''',(actor,imported))
                for sid, store in named.items():
                    amount = branch_values[sid]
                    if amount is None:
                        cur.execute('''update lfo.branch_target_plan set is_active=false,updated_at=now(),updated_by=%s
                            where distributor_id=%s and fy=%s and store_id=%s and is_active''',(actor,did,fy,sid))
                        continue
                    if run_id is None:
                        digest = hashlib.sha256(json.dumps([channel,fy,sorted(map(str,submitted))]).encode()).hexdigest()
                        cur.execute('''insert into lfo.import_run(source_id,sheet_name,content_hash,actor)
                            values(%s,%s,%s,%s) returning import_id''',(PLANNER_SOURCE,PLANNER_SHEET,digest,actor))
                        run_id = cur.fetchone()[0]
                    # A plan saved under another name for this store gives way first.
                    cur.execute('''update lfo.branch_target_plan set is_active=false,updated_at=now(),updated_by=%s
                        where distributor_id=%s and fy=%s and store_id=%s and is_active and normalized_name<>%s''',
                        (actor,did,fy,sid,store['key']))
                    cur.execute('''insert into lfo.branch_target_plan(distributor_id,fy,source_name,normalized_name,
                        monthly_target,yearly_target,store_id,is_active,source_id,sheet_name,
                        source_month_cell,source_year_cell,import_id,updated_by)
                        values(%s,%s,%s,%s,%s,%s,%s,true,%s,%s,%s,%s,%s,%s)
                        on conflict(distributor_id,fy,normalized_name) do update set
                        source_name=excluded.source_name,monthly_target=excluded.monthly_target,
                        yearly_target=excluded.yearly_target,store_id=excluded.store_id,is_active=true,
                        source_id=excluded.source_id,sheet_name=excluded.sheet_name,
                        source_month_cell=excluded.source_month_cell,source_year_cell=excluded.source_year_cell,
                        import_id=excluded.import_id,updated_by=excluded.updated_by,updated_at=now()''',
                        (did,fy,store['name'],store['key'],amount,amount*12,sid,PLANNER_SOURCE,PLANNER_SHEET,
                         PLANNER_SHEET,PLANNER_SHEET+' ×12',run_id,actor))
                    counts['branches'] += 1
            if run_id is not None:
                cur.execute('update lfo.import_run set counters=%s where import_id=%s',(Json(counts),run_id))
    return counts
